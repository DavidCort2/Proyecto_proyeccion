from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
import pandas as pd
import pytest

from core.config import PlanningRules
from core.curriculum import curriculum_key, schedule_name
from core.curriculum_planner import curriculum_coverage, execute_curriculum_plan, prepare_ficha_import
from core.curriculum_store import import_curricula, load_curricula
from core.database import load_planning, save_planning
from core.excel_parser import parse_instructors_excel
from core.export import export_planning
from core.fichas_parser import parse_fichas_excel

FIXTURES = Path(__file__).parent / "fixtures"
MALLA = FIXTURES / "orthotics" / "DESARROLLO Y ADAPTACION DE ORTESIS Y PROTESIS - DIURNO.xlsx"
PROGRAM = "DESARROLLO Y ADAPTACION DE PROTESIS Y ORTESIS"


@pytest.mark.parametrize("label,expected", [
    ("O&P-mañana", "Diurna"), ("P&O-tarde", "Diurna"), ("P & O - mañana", "Diurna"),
    ("Diurna O&P", "Diurna"), ("DiurnaO&P", "Diurna"), ("Diurno P&O tarde", "Diurna"),
    ("Mixta O&P", "Mixta"), ("O&P-Mixto", "Mixta"), ("Mixta", "Mixta"), ("DIURNO", "Diurna"),
])
def test_abbreviation_is_removed_before_reading_the_real_shift(label, expected):
    assert schedule_name(label) == expected
    assert curriculum_key(PROGRAM, label) == curriculum_key(PROGRAM, expected)


@pytest.mark.parametrize("label", ["O&P", "P&O", "O&P-nocturna"])
def test_abbreviation_without_a_supported_shift_does_not_invent_one(label):
    with pytest.raises(ValueError, match="Jornada no reconocida"):
        schedule_name(label)


def orthotics_rows():
    frame = parse_fichas_excel(Path(__file__).parents[1] / "data" / "reporteFichas_2026_4.xlsx")
    return frame.loc[frame["Especialidad"] == PROGRAM].copy()


def test_real_diurno_file_covers_all_three_fichas_without_an_extra_shift(tmp_path):
    path = tmp_path / "db.sqlite3"
    catalog = import_curricula(path, [(MALLA.name, MALLA.read_bytes())])
    staff = parse_instructors_excel(FIXTURES / "reporteInstructores.xlsx")
    frame = orthotics_rows()
    imported = prepare_ficha_import(frame, staff, catalog, 2027, "fichas.xlsx", "test", 2026, 4)
    assert len(imported["summary"]) == 1
    assert imported["summary"][0]["Jornada"] == "Diurna"
    assert imported["summary"][0]["Fichas que pasan"] == 3
    assert imported["summary"][0]["Fichas que terminan"] == 2
    assert {row["Jornada"] for row in imported["detail"]} == {"Diurna-tarde", "P&O-mañana", "P&O-tarde"}
    assert {row["Jornada de planeación"] for row in imported["detail"]} == {"Diurna"}
    assert {row["Archivo malla"] for row in imported["detail"]} == {MALLA.name}
    assert {row["Duración (trimestres)"] for row in imported["detail"]} == {10}
    assert {row["Trimestre actual"]: row["Fecha fin estimada"] for row in imported["detail"]} == {
        6: "2027-12-31", 8: "2027-06-30", 4: "2028-06-30"}
    coverage = curriculum_coverage(pd.DataFrame(imported["summary"]), catalog)
    assert coverage["Estado"].tolist() == ["Lista"]
    assert coverage["Archivo malla"].tolist() == [MALLA.name]
    plan = execute_curriculum_plan(staff, imported, catalog, PlanningRules(), {"Técnico": 0, "Tecnólogo": 0}, 2027, "planta.xlsx", "test")
    assert [row["Horas requeridas"] for row in plan["quarterly"]] == [1404, 1440, 960, 960]
    assert [row["Fichas activas"] for row in plan["quarterly"]] == [3, 3, 2, 2]
    assert plan["center"]["demanda_total_horas_anuales"] == 4764
    assert sum(row["Horas asignadas (h/mes)"] for row in plan["monthly_assignments"]) == pytest.approx(4764)
    assert len(plan["distribution"]) == 1
    assert {row["Jornada"] for row in plan["monthly_fichas"]} == {"Diurna"}
    assert not pd.DataFrame(plan["monthly_fichas"]).duplicated(["Ficha", "Mes número"]).any()
    assert len(load_curricula(path)["curricula"]) == 1
    save_planning(path, staff, plan)
    book = load_workbook(BytesIO(export_planning(*load_planning(path))), data_only=True)
    sheet = book["Distribucion"]
    columns = [cell.value for cell in sheet[1]]
    assert [row[columns.index("Jornada")] for row in sheet.iter_rows(min_row=2, values_only=True)] == ["Diurna"]


def test_new_fichas_use_the_same_real_curriculum_and_real_shift(tmp_path):
    catalog = import_curricula(tmp_path / "db.sqlite3", [(MALLA.name, MALLA.read_bytes())])
    staff = parse_instructors_excel(FIXTURES / "reporteInstructores.xlsx")
    frame = orthotics_rows().loc[lambda row: row["Jornada"].eq("P&O-mañana")].copy()
    frame["Trimestre actual"] = catalog["curricula"][0]["duration"]
    imported = prepare_ficha_import(frame, staff, catalog, 2027, "fichas.xlsx", "test", 2026, 4)
    plan = execute_curriculum_plan(staff, imported, catalog, PlanningRules(intake_weights=(100, 0, 0, 0)),
                                   {"Técnico": 0, "Tecnólogo": 25}, 2027, "planta.xlsx", "test")
    assert [row["Horas requeridas"] for row in plan["quarterly"]] == [396, 480, 480, 480]
    assert plan["center"]["demanda_total_horas_anuales"] == 1836
    assert {row["Jornada"] for row in plan["monthly_fichas"]} == {"Diurna"}
    assert {row["Fecha fin estimada"] for row in plan["monthly_fichas"]} == {"2029-06-30"}
