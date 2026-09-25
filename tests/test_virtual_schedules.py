from copy import deepcopy
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from core.curriculum_store import import_curricula
from core.database import load_planning, reset_planning_database, save_planning
from core.export import export_planning
from core.virtual_competencies import competency_rows
from core.virtual_schedule import duration_boundary, lective_activities, parse_virtual_schedule
from core.virtual_schedule_store import import_virtual_schedules, load_virtual_schedules, save_virtual_competencies
from core.virtual_schedule_planner import VirtualRules, execute_virtual_schedule_plan
from core.virtual_staffing import workdays
from test_curriculum_automation import malla

FIXTURES = Path(__file__).parent / "fixtures/virtual_schedules"
REAL_SCHEDULE = FIXTURES / "cronograma_adso.xlsx"


def schedule_bytes(program="Software", durations=(7, 7, 7), start=date(2026, 1, 1), productive_days=0, transversal_code="240202501"):
    book = Workbook()
    sheet = book.active
    sheet.title = "Cronograma"
    end = start + timedelta(days=sum(durations) + productive_days - 1)
    sheet.append(["CRONOGRAMA GENERAL DE FORMACIÓN TITULADA VIRTUAL"])
    sheet.append(["Nombre de programa: " + program])
    sheet.append(["Fecha Inicio:", None, None, None, start])
    sheet.append(["Fecha Fin:", None, None, None, end])
    sheet.append(["Fases", "Actividades del proyecto", "Actividades de Aprendizaje", "Tiempo de duración estimado Días",
                  "Tiempo de duración estimado Horas", "Fecha Inicio", "Fecha Final"])
    for i, days in enumerate(durations, 1):
        finish = start + timedelta(days=days - 1)
        code = transversal_code if i == 2 else "220501092"
        text = "Conversación en inglés" if i == 2 else "Desarrollo de software"
        sheet.append([f"Fase {i} {'Análisis' if i == 1 else 'Planeación' if i == 2 else 'Ejecución'}", f"AP{i} Proyecto",
                      f"GA{i}-{code}-AA1 {text}", days, f"{i * 100} HORAS", start, finish])
        start = finish + timedelta(days=1)
    if productive_days:
        sheet.append(["ETAPA PRODUCTIVA", None, None, productive_days, "864 HORAS", start, end])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def configured_catalog(tmp_path, **kwargs):
    path = tmp_path / "virtual.sqlite3"
    return path, import_virtual_schedules(path, [("Cronograma.xlsx", schedule_bytes(**kwargs))])


def plan(catalog, **kwargs):
    arguments = dict(programs=[{"Programa": "Software", "Incluir": True, "Nivel": "Tecnólogo", "Peso de oferta": 1}],
                     cohorts=[], plant=[], rules=VirtualRules(intake_weights=(100, 0, 0, 0)),
                     targets={"Técnico": 0, "Tecnólogo": 25}, year=2027)
    arguments.update(kwargs)
    return execute_virtual_schedule_plan(catalog, **arguments)


def person(document="00123", kind="Técnico", profile="Software"):
    return {"Nombre completo": "Ana María Pérez", "Cédula": document, "Tipo": kind, "Perfil": profile}


def save_choices(path, catalog, choices):
    return save_virtual_competencies(path, {item["program_key"]: item["source_digest"] for item in catalog["schedules"]}, choices)


def real_catalog(tmp_path):
    return import_virtual_schedules(tmp_path / "real.sqlite3", [(name, (FIXTURES / name).read_bytes())
                                                               for name in ("adso_fases.xlsx", "ciberseguridad_fases.xlsx")])


def test_users_two_files_have_complete_phases_duration_and_unique_competencies(tmp_path):
    catalog = real_catalog(tmp_path)
    adso, cyber = catalog["schedules"]
    assert adso["lective_duration"] == 21
    assert adso["blocks"][-1]["duration"] == 2.75
    assert adso["duration_adjustments"][0]["added_duration"] == 1.75
    assert len(adso["activities"]) == 98
    assert len(cyber["activities"]) == 56
    assert cyber["lective_duration"] == 9
    assert {a["competency"] for a in cyber["activities"] if a["source_row"] == 31} == {"240201528"}
    assert cyber["blocks"][2]["source_notes"][0]["source_row"] == 32
    assert "relacionadas con expresiones" not in next(a["activity"] for a in cyber["activities"] if a["source_row"] == 31)
    assert not any(a["source_row"] == 32 for a in cyber["activities"])
    unique = competency_rows(catalog)
    assert len(unique) == 24  # 23 códigos distintos y la inducción común.
    assert len([row for row in unique if row["Competencia"] == "240202501"]) == 1
    assert len({a["competency"] for a in adso["activities"]}) == 19
    assert len({a["competency"] for a in cyber["activities"]}) == 15
    # Filas sin proyecto/duración ubicadas antes del encabezado de su bloque.
    adso_rows = {a["source_row"]: a for a in adso["activities"]}
    assert adso_rows[32]["block_id"] == adso_rows[33]["block_id"]
    assert adso_rows[32]["phase"] == "Fase 2 · Planeación"
    assert {adso_rows[n]["block_id"] for n in (61, 62, 63, 64)} == {adso_rows[64]["block_id"]}
    assert adso_rows[61]["phase"] == "Fase 3 · Ejecución"
    cyber_rows = {a["source_row"]: a for a in cyber["activities"]}
    assert cyber_rows[34]["block_id"] == cyber_rows[35]["block_id"]
    assert cyber_rows[34]["phase"] == "Fase 3 · Ejecución"
    assert {row["Tipo"] for row in unique if row["Competencia"] == "240201526"} == {"Transversal"}
    for item in catalog["schedules"]:
        assert item["schema_version"] == 3
        assert "reference_start" not in item and "reference_end" not in item
        assert all("start" not in row and "end" not in row for row in item["activities"])


def test_legacy_workbook_merged_blocks_are_read_without_double_counting():
    item = parse_virtual_schedule(REAL_SCHEDULE.read_bytes(), REAL_SCHEDULE.name)
    assert len(item["blocks"]) == 11
    assert sum(row["source_hours"] for row in item["blocks"] if row["phase"].startswith("Fase")) == 3072
    assert item["blocks"][-1]["source_hours"] == 864
    assert item["activities"][-1]["activity"] == "Seguimiento de etapa productiva"
    assert {row["competency"] for row in item["activities"] if row["source_row"] == 43} == {"220501093", "220501095"}
    assert all(row["teaching_type"] in {"Técnico", "Transversal"} for row in item["activities"])


def test_source_dates_and_manual_hours_never_affect_demand(tmp_path):
    path, catalog = configured_catalog(tmp_path)
    _, before = plan(catalog)
    changed = import_virtual_schedules(path, [("Cronograma.xlsx", schedule_bytes(start=date(1999, 5, 7)))])
    for activity in changed["schedules"][0]["activities"]:
        activity["instructor_hours"] = 999999
    _, after = plan(changed)
    for key in ("cohort_dates", "periods", "contracts", "monthly", "center", "summary"):
        assert before[key] == after[key]
    assert after["center"]["demanda_total_horas_anuales"] == 22  # Dos semanas técnicas de 10 h y una transversal de 2 h.


def test_broken_or_missing_dates_do_not_block_template(tmp_path):
    book = load_workbook(BytesIO(schedule_bytes()))
    sheet = book.active
    for row in sheet:
        for cell in row:
            if cell.column >= 6 or cell.coordinate in {"E3", "E4"}:
                cell.value = "NO ES UNA FECHA"
    output = BytesIO()
    book.save(output)
    item = parse_virtual_schedule(output.getvalue(), "Sin fechas.xlsx")
    assert item["lective_duration"] == 21 and item["duration_unit"] == "días"


def test_missing_duration_rejected_instead_of_inferred_from_cohort_dates():
    book = load_workbook(BytesIO(schedule_bytes()))
    book.active["D5"] = "Otra columna"
    output = BytesIO()
    book.save(output)
    with pytest.raises(ValueError, match="no se deduce de las fechas"):
        parse_virtual_schedule(output.getvalue(), "Sin duración.xlsx")


def test_manual_classification_applies_across_programs_phases_and_reimports(tmp_path):
    path, catalog = configured_catalog(tmp_path, transversal_code="123456789")
    catalog = import_virtual_schedules(path, [("Redes.xlsx", schedule_bytes(program="Redes", transversal_code="123456789"))])
    rows = competency_rows(catalog)
    for row in rows:
        row["Tipo"] = "Técnico"
    save_choices(path, catalog, rows)
    catalog = import_virtual_schedules(path, [("Redes.xlsx", schedule_bytes(program="Redes", durations=(14, 7, 7), transversal_code="123456789"))])
    assert all(row["Tipo"] == "Técnico" for row in competency_rows(catalog))
    assert all(a["classification_source"] == "manual" for i in catalog["schedules"] for a in i["activities"] if a["competency"] == "123456789")


def test_continuing_end_date_reconstructs_only_pending_phases(tmp_path):
    _, catalog = configured_catalog(tmp_path)
    _, execution = plan(catalog, targets={"Técnico": 0, "Tecnólogo": 100}, cohorts=[
        {"Programa": "Software", "Fichas que pasan": 2, "Fecha fin lectiva": "2027-01-14"}])
    assert execution["center"]["fichas_nuevas"] == 2
    assert execution["center"]["demanda_total_horas_anuales"] == 68  # 2 × 12 pendientes + 2 × 22 nuevas.
    assert execution["cohort_dates"][0]["Fases al iniciar la vigencia"] == "Fase 2 · Planeación"
    assert execution["cohort_dates"][0]["Fecha fin lectiva"] == "2027-01-14"
    assert all(row["Mes"] == 1 for row in execution["monthly_fichas"])


def test_variable_phase_lengths_and_end_of_contracts(tmp_path):
    _, catalog = configured_catalog(tmp_path, durations=(7, 14, 7))
    _, execution = plan(catalog, plant=[person(kind="Transversal", profile="240202501")])
    assert execution["center"]["demanda_total_horas_anuales"] == 24
    assert execution["summary"]["pico_contratistas_total"] == 1
    assert [(r["Inicio"], r["Fin"]) for r in execution["contracts"]] == [("2027-01-01", "2027-01-07"), ("2027-01-22", "2027-01-28")]
    assert execution["cohort_dates"][0]["Fecha fin lectiva"] == "2027-01-28"


def test_offers_are_automatic_and_rounded_once_from_remaining_meta(tmp_path):
    _, catalog = configured_catalog(tmp_path)
    _, execution = plan(catalog, targets={"Técnico": 0, "Tecnólogo": 101}, rules=VirtualRules(intake_weights=(25, 25, 25, 25)))
    assert execution["center"]["fichas_nuevas"] == 5
    assert execution["virtual_inputs"]["offers"] == ["2027-01-01", "2027-04-01", "2027-07-01", "2027-10-01"]
    assert [row["Fecha inicio lectiva estimada"] for row in execution["cohort_dates"]] == execution["virtual_inputs"]["offers"]
    assert [execution["offers_by_program"][0][f"Oferta {i}"] for i in range(1, 5)] == [2, 1, 1, 1]
    assert execution["center"]["demanda_total_horas_anuales"] == 110


@pytest.mark.parametrize("fichas,plants,required", [(4, 0, 1), (5, 0, 2), (3, 1, 0), (4, 1, 1), (6, 2, 0), (7, 2, 1)])
def test_daily_capacity_uses_whole_fichas_not_pooled_spare_hours(tmp_path, fichas, plants, required):
    _, catalog = configured_catalog(tmp_path, durations=(7,))
    _, execution = plan(catalog, targets={"Técnico": 0, "Tecnólogo": fichas * 25}, plant=[person(str(i)) for i in range(plants)])
    assert execution["summary"]["pico_contratistas_total"] == required
    assert sum(row["Horas asignadas"] for row in execution["monthly_assignments"]) == fichas * 10
    for row in execution["monthly_instructors"]:
        assert row["Pico fichas asignadas"] <= (3 if row["Vinculación"] == "Planta" else 4)
        assert 0 <= row["Horas asignadas"] <= row["Capacidad en horas"]


def test_transversal_capacity_shared_once_between_programs(tmp_path):
    path, catalog = configured_catalog(tmp_path)
    catalog = import_virtual_schedules(path, [("Redes.xlsx", schedule_bytes(program="Redes"))])
    staff, execution = plan(catalog, targets={"Técnico": 0, "Tecnólogo": 100},
                            programs=[{"Programa": name, "Incluir": True, "Nivel": "Tecnólogo", "Peso de oferta": 1}
                                      for name in ("Software", "Redes")],
                            plant=[person(kind="Transversal", profile="240202501")])
    shared = next(row for row in execution["staffing"] if row["Tipo"] == "Transversal" and row["Inicio"] == "2027-01-08")
    assert len(staff) == 1
    assert shared["Horas requeridas (h/sem)"] == 8
    assert shared["Capacidad planta (h/sem)"] == 32
    assert shared["Horas a contratar"] == 0
    assert shared["Contratistas requeridos"] == 0
    assert shared["Fichas cubiertas por planta"] == 4


def test_multiple_activities_do_not_multiply_technical_or_transversal_load(tmp_path):
    _, catalog = configured_catalog(tmp_path)
    _, original = plan(catalog)
    item = catalog["schedules"][0]
    for row in deepcopy(item["activities"]):
        row["id"] += "duplicate"
        row["activity"] += " Otra evidencia"
        if row["teaching_type"] == "Técnico":
            row["competency"] = "220501999"  # Otro contenido del mismo perfil técnico.
        item["activities"].append(row)
    _, changed = plan(catalog)
    assert changed["monthly"] == original["monthly"]
    assert changed["contracts"] == original["contracts"]


@pytest.mark.parametrize("productive_hours", [None, 0, 1000])
def test_productive_stage_never_contributes_hours_or_dates(tmp_path, productive_hours):
    _, catalog = configured_catalog(tmp_path, productive_days=180)
    catalog["schedules"][0]["activities"][-1].update(teaching_type="Transversal", instructor_hours=productive_hours)
    _, execution = plan(catalog)
    assert execution["center"]["demanda_total_horas_anuales"] == 22
    assert execution["cohort_dates"][0]["Fecha fin lectiva"] == "2027-01-21"
    assert all(row["Fase"] != "Etapa productiva" for row in execution["activity_hours"])
    assert all(row["Fin"] <= "2027-01-21" for row in execution["contracts"])


@pytest.mark.parametrize("kwargs", [
    {"plant": [person(kind="Técnico", profile="240202501")]},
    {"plant": [{**person(), "Nombre completo": ""}]},
    {"plant": [{**person(), "Cédula": ""}]},
    {"plant": [{**person(), "Cédula": 12.5}]},
    {"plant": [person(), person(kind="Transversal", profile="240202501")]},
    {"cohorts": [{"Programa": "Software", "Fichas que pasan": 1, "Fecha fin lectiva": "2026-12-31"}]},
    {"cohorts": [{"Programa": "Software", "Fichas que pasan": 1, "Fecha fin lectiva": "2027-12-31"}]},
    {"rules": VirtualRules(intake_weights=(10, 10, 10, 10))},
    {"rules": VirtualRules(weekly_plant_direct_hours=9)},
])
def test_invalid_inputs_do_not_produce_a_plan(tmp_path, kwargs):
    _, catalog = configured_catalog(tmp_path)
    with pytest.raises(ValueError):
        plan(catalog, **kwargs)


def test_import_atomicity_and_stale_classification(tmp_path):
    path, before = configured_catalog(tmp_path)
    with pytest.raises(ValueError):
        import_virtual_schedules(path, [("Redes.xlsx", schedule_bytes(program="Redes")), ("Otro.xlsx", b"invalid")])
    assert load_virtual_schedules(path) == before
    with pytest.raises(ValueError, match="cambiaron"):
        save_virtual_competencies(path, {}, competency_rows(before))
    with pytest.raises(ValueError, match="una sola vez"):
        save_choices(path, before, competency_rows(before) * 2)
    assert load_virtual_schedules(path) == before


def test_old_catalog_requires_reupload_not_inferred_duration(tmp_path):
    _, catalog = configured_catalog(tmp_path)
    catalog["schedules"][0]["schema_version"] = 2
    with pytest.raises(ValueError, match="vuelva a cargar"):
        plan(catalog)


def test_real_files_monthly_hours_assignments_and_exports_reconcile(tmp_path):
    catalog = real_catalog(tmp_path)
    adso, cyber = catalog["schedules"]
    staff, execution = plan(catalog, targets={"Técnico": 100, "Tecnólogo": 200},
        programs=[{"Programa": i["program"], "Incluir": True, "Nivel": level, "Peso de oferta": 1}
                  for i, level in [(adso, "Tecnólogo"), (cyber, "Técnico")]],
        cohorts=[{"Programa": adso["program"], "Fichas que pasan": 2, "Fecha fin lectiva": "2027-04-30"},
                 {"Programa": cyber["program"], "Fichas que pasan": 1, "Fecha fin lectiva": "2027-02-28"}],
        plant=[person(profile=adso["program"]), person("00456", "Transversal", "240202501")], rules=VirtualRules())
    assert execution["center"]["fichas_que_pasan"] == 3
    assert execution["center"]["fichas_nuevas"] == 9
    assert execution["cohort_dates"][0]["Fecha fin lectiva"] == "2027-04-30"
    for month in execution["monthly"]:
        number = month["Mes"]
        assert month["Horas requeridas"] == pytest.approx(sum(r["Horas requeridas"] for r in execution["monthly_fichas"] if r["Mes"] == number))
        assert month["Horas requeridas"] == pytest.approx(sum(r["Horas asignadas"] for r in execution["monthly_assignments"] if r["Mes"] == number))
        assert month["Horas requeridas"] == pytest.approx(month["Horas cubiertas por planta"] + month["Horas a contratar"])
        assert all(r["Horas disponibles"] >= 0 for r in execution["monthly_instructors"] if r["Mes"] == number)
    assert sum(row["Horas instructor en vigencia"] for row in execution["activity_hours"]) == pytest.approx(execution["center"]["demanda_total_horas_anuales"])
    continuing = [row for row in execution["monthly_fichas"] if row["Cohorte"] == "Continuación 1"]
    assert max(row["Mes"] for row in continuing) == 4
    assert "Evaluación" in continuing[-1]["Fases"]
    path = tmp_path / "output.sqlite3"
    save_planning(path, staff, execution)
    restored = load_planning(path)
    assert restored[0].iloc[0]["Documento"] == "00123"
    book = load_workbook(BytesIO(export_planning(*restored)), read_only=True)
    assert {"Competencias unicas", "Horas mensuales por ficha", "Capacidad por instructor", "Resumen trimestral"} <= set(book.sheetnames)
    assert book["Competencias unicas"].max_row == 25
    book.close()


def test_export_save_reset_and_calculation_do_not_touch_presencial(tmp_path):
    path, catalog = configured_catalog(tmp_path)
    presencial = tmp_path / "presencial.sqlite3"
    import_curricula(presencial, [("Software - DIURNA.xlsx", malla([12, 12]))])
    snapshot = presencial.read_bytes()
    staff, execution = plan(catalog)
    save_planning(path, staff, execution)
    book = load_workbook(BytesIO(export_planning(*load_planning(path))), read_only=True)
    assert "Contratistas y fechas" in book.sheetnames
    book.close()
    reset_planning_database(path)
    assert load_virtual_schedules(path) == {"schedules": []}
    assert load_planning(path) is None
    assert presencial.read_bytes() == snapshot


@pytest.mark.parametrize("start,offset,unit,expected", [
    (date(2027, 1, 1), 21, "meses", date(2028, 10, 1)),
    (date(2027, 1, 1), 0.25, "meses", date(2027, 1, 9)),
    (date(2028, 1, 31), 1, "meses", date(2028, 2, 29)),
    (date(2027, 5, 1), -2.75, "meses", date(2027, 2, 8)),
    (date(2027, 1, 1), -3, "semanas", date(2026, 12, 11)),
])
def test_calendar_duration_boundaries(start, offset, unit, expected):
    assert duration_boundary(start, offset, unit) == expected


def test_weekends_and_empty_demand_have_no_contracts(tmp_path):
    _, catalog = configured_catalog(tmp_path)
    _, execution = plan(catalog, targets={"Técnico": 0, "Tecnólogo": 0}, plant=[person()])
    assert execution["summary"]["pico_contratistas_total"] == 0
    assert execution["contracts"] == []
    assert execution["center"]["demanda_total_horas_anuales"] == 0
    assert sum(r["Capacidad en horas"] for r in execution["monthly_instructors"]) == workdays(date(2027, 1, 1), date(2028, 1, 1)) * 32 / 5


def test_decimal_months_do_not_add_a_day_or_overlap_phases(tmp_path):
    book = load_workbook(BytesIO(schedule_bytes()))
    sheet = book.active
    sheet["D5"] = "Tiempo de duración estimado Meses"
    for row, duration in zip((6, 7, 8), (0.1, 0.2, 0.7)):
        sheet.cell(row, 4).value = duration
    content = BytesIO()
    book.save(content)
    catalog = import_virtual_schedules(tmp_path / "months.sqlite3", [("Meses.xlsx", content.getvalue())])
    assert [block["start_offset"] for block in catalog["schedules"][0]["blocks"]] == [0, 0.1, 0.3]
    _, result = plan(catalog, rules=VirtualRules(intake_weights=(0, 100, 0, 0)))
    intervals = result["activity_hours"]
    assert intervals[1]["Fin"] == "2027-04-09"
    assert intervals[2]["Inicio"] == "2027-04-10"
    assert result["center"]["demanda_total_horas_anuales"] == 36  # 17 días técnicos y una semana transversal.


def test_two_active_transversals_each_add_one_profile_load(tmp_path):
    _, catalog = configured_catalog(tmp_path, durations=(7,))
    activities = catalog["schedules"][0]["activities"]
    for code in ("240202501", "240201526"):
        activity = deepcopy(activities[0])
        activity.update(id=code, competency=code, teaching_type="Transversal")
        activities.append(activity)
    _, execution = plan(catalog, targets={"Técnico": 0, "Tecnólogo": 100})
    assert execution["summary"]["pico_contratistas_total"] == 2
    assert execution["summary"]["tecnicos_en_pico"] == 1
    assert execution["summary"]["transversales_en_pico"] == 1
    assert execution["center"]["demanda_total_horas_anuales"] == 56  # 4 fichas × (10 técnicas + 2 + 2 transversales).


def test_partial_year_and_reduced_later_demand(tmp_path):
    _, catalog = configured_catalog(tmp_path, durations=(365,))
    _, execution = plan(catalog, targets={"Técnico": 0, "Tecnólogo": 300},
        cohorts=[{"Programa": "Software", "Fichas que pasan": 8, "Fecha fin lectiva": "2027-03-31"}],
        rules=VirtualRules(intake_weights=(0, 0, 0, 100)))
    assert execution["summary"]["pico_contratistas_total"] == 2
    assert execution["summary"]["trimestre_pico"] == 1
    assert [r["Pico simultáneo de contratistas"] for r in execution["quarterly"]] == [2, 0, 0, 1]
    assert [r["Cupos excedentes si se mantiene el pico anual"] for r in execution["quarterly"]] == [0, 2, 2, 1]
    assert [(r["Inicio"], r["Fin"]) for r in execution["contracts"]] == [
        ("2027-01-01", "2027-03-31"), ("2027-10-01", "2027-12-31"), ("2027-01-01", "2027-03-31")]
    assert execution["contracts"][1]["Fin limitado por vigencia"]
    expected = 8 * 2 * workdays(date(2027, 1, 1), date(2027, 4, 1)) + 4 * 2 * workdays(date(2027, 10, 1), date(2028, 1, 1))
    assert execution["center"]["demanda_total_horas_anuales"] == expected
