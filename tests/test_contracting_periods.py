from io import BytesIO

from openpyxl import Workbook, load_workbook
import pytest

from core.config import PlanningRules
from core.contracting_periods import contracting_windows
from core.curriculum_planner import execute_curriculum_plan, prepare_ficha_import
from core.curriculum_store import import_curricula
from core.database import load_planning, save_planning
from core.export import export_planning
from test_curriculum import PROGRAM, ficha_frame, instructors_frame, run_plan
from test_curriculum_automation import malla


@pytest.mark.parametrize("counts,expected", [
    ([27, 21, 10, 0], [(1, 1, 6), (1, 2, 11), (1, 3, 10)]),
    ([3, 1, 3, 1], [(1, 1, 2), (1, 4, 1), (3, 3, 2)]),
    ([1, 2, 3, 4], [(1, 4, 1), (2, 4, 1), (3, 4, 1), (4, 4, 1)]),
    ([0, 0, 0, 0], []),
])
def test_contracts_follow_demand_and_do_not_cover_unneeded_quarters(counts, expected):
    assert contracting_windows(counts) == expected


def changing_curriculum():
    book = Workbook()
    book.remove(book.active)
    for quarter, transversal in enumerate([40, 20, 4, 0], 1):
        sheet = book.create_sheet(f"Trimestre {quarter}")
        sheet.append(["Competencia", "Resultado", "Horas semanales", "Tipo competencia"])
        sheet.append(["Competencia técnica", "Resultado técnico", 40 - transversal, "Técnica"])
        sheet.append(["Competencia transversal", "Resultado transversal", transversal, "Integralidad"])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_transversal_peak_reduces_as_same_fichas_advance_and_profiles_peak_separately(tmp_path):
    catalog = import_curricula(tmp_path / "db.sqlite3", [(PROGRAM + " - DIURNA.xlsx", changing_curriculum())])
    _, plan = run_plan(catalog, frame=ficha_frame(day_age=4).iloc[:1], target=675)
    quarters = plan["contracting_quarterly"]
    assert [row["Fichas activas"] for row in quarters] == [27] * 4
    assert [row["Contratistas transversales"] for row in quarters] == [27, 13, 2, 0]
    assert [row["Contratistas técnicos"] for row in quarters] == [0, 13, 24, 27]
    assert [row["Exceso transversal si conserva picos por perfil"] for row in quarters] == [0, 14, 25, 27]
    assert plan["summary"]["pico_contratistas_total"] == 27  # No suma los dos picos separados (54).
    assert plan["summary"]["trimestres_pico_total"] == [1, 4]
    assert plan["summary"]["trimestres_pico_transversal"] == [1]
    windows = [row for row in plan["contract_windows"] if row["Área"] == "Integralidad"]
    assert [(row["Contratistas"], row["Inicio"], row["Fin"]) for row in windows] == [
        (14, "2027-01-01", "2027-03-31"), (11, "2027-01-01", "2027-06-30"), (2, "2027-01-01", "2027-09-30")]
    integral = [row for row in plan["contracting_profile_quarterly"] if row["Área"] == "Integralidad"]
    assert [row["Contratos que finalizan"] for row in integral] == [14, 11, 2, 0]
    assert [row["Reducción respecto a T anterior"] for row in integral] == [0, 14, 11, 2]
    assert all(row["Horas asignadas (h/mes)"] <= row["Capacidad (h/mes)"] for row in plan["monthly_instructors"])


@pytest.mark.parametrize("offer,annual", [(1, 1200), (2, 720), (3, 360), (4, 120)])
def test_every_offer_starts_at_training_quarter_one_and_uses_only_remaining_calendar(offer, annual, tmp_path):
    catalog = import_curricula(tmp_path / "db.sqlite3", [(PROGRAM + " - DIURNA.xlsx", malla([10, 20, 30, 40]))])
    rules = PlanningRules(intake_weights=tuple(100 if q == offer else 0 for q in range(1, 5)))
    _, plan = run_plan(catalog, frame=ficha_frame(day_age=4).iloc[:1], target=25, rules=rules)
    assert plan["center"]["demanda_total_horas_anuales"] == annual
    assert {row["Mes número"] for row in plan["monthly_fichas"]} == set(range(3 * offer - 2, 13))
    for row in plan["monthly_fichas"]:
        assert row["Trimestre de formación"] == row["Trimestre"] - offer + 1
        assert row["Horas requeridas (h/mes)"] == (row["Trimestre"] - offer + 1) * 40


def test_current_contractors_cannot_change_hours_staffing_periods_or_assignments(tmp_path):
    catalog = import_curricula(tmp_path / "db.sqlite3", [(PROGRAM + " - DIURNA.xlsx", changing_curriculum())])
    frame = ficha_frame(day_age=4).iloc[:1]
    staff, baseline = run_plan(catalog, frame=frame, target=675)
    for index, (area, profile) in enumerate([("Técnica", PROGRAM), ("Bilingüismo", "Bilingüismo"), ("Integralidad", "Integralidad")], 10):
        staff.loc[len(staff)] = [profile, area, f"Contrato existente {index}", str(index), "Contratista", 9999, False]
    source = prepare_ficha_import(frame, staff, catalog, 2027, "fichas.xlsx", "fichas", 2026, 4)
    plan = execute_curriculum_plan(staff, source, catalog, PlanningRules(intake_weights=(100, 0, 0, 0)),
                                   {"Técnico": 0, "Tecnólogo": 675}, 2027, "instructores.xlsx", "instructores")
    for key in ("monthly", "monthly_staffing", "monthly_instructors", "monthly_assignments", "contract_windows", "contracting_quarterly", "summary"):
        assert plan[key] == baseline[key]
    assert "transversal_current" not in plan
    assert "transversal_continuity" not in plan
    save_planning(tmp_path / "plan.sqlite3", staff, plan)
    book = load_workbook(BytesIO(export_planning(*load_planning(tmp_path / "plan.sqlite3"))), data_only=True)
    exported = " ".join(str(cell) for sheet in book for row in sheet.values for cell in row)
    assert "Contrato existente" not in exported
    assert "Contratistas actuales" not in exported
    assert "Contratistas adicionales" not in exported


def test_without_plant_entire_curricular_demand_is_contracted(tmp_path):
    catalog = import_curricula(tmp_path / "db.sqlite3", [(PROGRAM + " - DIURNA.xlsx", malla([40, 40, 40, 40]))])
    staff = instructors_frame()
    staff["Es planta"] = False
    staff["Tipo Contrato"] = "Contratista"
    source = prepare_ficha_import(ficha_frame(day_age=4).iloc[:1], staff, catalog, 2027, "f.xlsx", "x", 2026, 4)
    plan = execute_curriculum_plan(staff, source, catalog, PlanningRules(intake_weights=(100, 0, 0, 0)),
                                   {"Técnico": 0, "Tecnólogo": 50}, 2027, "i.xlsx", "x")
    assert all(row["Contratistas requeridos"] == 2 for row in plan["monthly"])
    assert all(row["Horas a contratar"] == row["Horas requeridas"] == 320 for row in plan["monthly"])
    assert all(row["Tipo"] == "Contratista proyectado" for row in plan["monthly_instructors"])


def test_sufficient_plant_needs_no_contracting_periods(tmp_path):
    catalog = import_curricula(tmp_path / "db.sqlite3", [(PROGRAM + " - DIURNA.xlsx", malla([5, 5, 5, 5]))])
    staff, plan = run_plan(catalog, frame=ficha_frame(day_age=4).iloc[:1], target=25)
    assert plan["contract_windows"] == []
    assert plan["summary"]["pico_contratistas_total"] == 0
    assert plan["summary"]["trimestres_pico_total"] == []
    assert plan["summary"]["meses_contratista_requeridos"] == 0
    assert all(row["Tipo"] == "Planta" for row in plan["monthly_instructors"])
    save_planning(tmp_path / "plan.sqlite3", staff, plan)
    assert export_planning(*load_planning(tmp_path / "plan.sqlite3"))


def test_decimal_curricular_hours_do_not_create_a_phantom_contractor(tmp_path):
    book = Workbook()
    book.remove(book.active)
    for q in range(1, 5):
        sheet = book.create_sheet(f"Trimestre {q}")
        sheet.append(["Competencia", "Resultado", "Horas semanales", "Tipo competencia"])
        for result in range(20):
            sheet.append(["Competencia técnica", f"Resultado {result}", 0.1, "Técnica"])
    stream = BytesIO()
    book.save(stream)
    catalog = import_curricula(tmp_path / "db.sqlite3", [(PROGRAM + " - DIURNA.xlsx", stream.getvalue())])
    rules = PlanningRules(weekly_plant_direct_hours=1, weekly_contractor_hours=1, intake_weights=(100, 0, 0, 0))
    _, plan = run_plan(catalog, frame=ficha_frame(day_age=4).iloc[:1], target=25, rules=rules)
    # Veinte resultados de 0,1 h demandan 2 h: planta 1 h + un contratista 1 h.
    assert all(row["contratistas_totales"] == 1 for row in plan["quarterly"])
    assert all(row["Contratistas requeridos"] == 1 for row in plan["monthly"])
    assert plan["summary"]["contratistas_totales"] == plan["summary"]["pico_contratistas_total"] == 1
