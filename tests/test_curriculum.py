from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from core.config import PlanningRules
from core.curriculum import name_key, parse_curriculum
from core.curriculum_store import import_curricula, load_curricula, save_competencies
from core.curriculum_planner import execute_curriculum_plan, prepare_ficha_import
from core.database import load_planning, save_planning
from core.excel_parser import parse_instructors_excel
from core.export import export_planning
from core.fichas_parser import parse_fichas_excel

FIXTURES = Path(__file__).parent / "fixtures" / "curricula"
PROGRAM = "ANALISIS Y DESARROLLO DE SOFTWARE"


@pytest.fixture
def curriculum_files():
    return [(p.name, p.read_bytes()) for p in sorted(FIXTURES.glob("*.xlsx"))]


@pytest.fixture
def curriculum_catalog(tmp_path, curriculum_files):
    return import_curricula(tmp_path / "curricula.sqlite3", curriculum_files)


def instructors_frame():
    # Un solo perfil técnico compartido por jornadas y una planta por transversal.
    frame = pd.DataFrame([
        [PROGRAM, "Técnica", "Planta técnica", "01", "Planta", 32.0, True],
        ["Bilingüismo", "Bilingüismo", "Planta inglés", "02", "Planta", 32.0, True],
        ["Integralidad", "Integralidad", "Planta integral", "03", "Planta", 32.0, True],
    ], columns=["Especialidad", "Área", "Nombre", "Documento", "Tipo Contrato", "Horas programadas actuales", "Es planta"])
    frame.attrs["specialties"] = frame[["Especialidad", "Área"]].to_dict("records")
    return frame


def ficha_frame(day_age=6, mixed_age=8):
    return pd.DataFrame([
        {"Ficha": "001", "Especialidad": PROGRAM, "Nivel": "Tecnólogo", "Jornada": "Diurna-mañana", "Trimestre actual": day_age},
        {"Ficha": "002", "Especialidad": PROGRAM, "Nivel": "Tecnólogo", "Jornada": "Mixta", "Trimestre actual": mixed_age},
    ])


def run_plan(catalog, frame=None, rules=None, target=100):
    instructors = instructors_frame()
    imported = prepare_ficha_import(ficha_frame() if frame is None else frame, instructors, catalog, 2027,
                                    "fichas.xlsx", "fichas", 2026, 4)
    execution = execute_curriculum_plan(instructors, imported, catalog, rules or PlanningRules(intake_weights=(100, 0, 0, 0)),
                                        {"Técnico": 0, "Tecnólogo": target}, 2027, "instructores.xlsx", "instructores")
    return instructors, execution


def test_real_workbooks_are_complete_and_preserve_outcomes(curriculum_catalog):
    assert len(curriculum_catalog["curricula"]) == 2
    assert len(curriculum_catalog["competencies"]) == 18
    assert sum(c["transversal"] for c in curriculum_catalog["competencies"]) == 11
    expected = {"Diurna": (7, 44, 30, 2520), "Mixta": (9, 45, 26, 2808)}
    for item in curriculum_catalog["curricula"]:
        duration, rows, weekly, total = expected[item["schedule"]]
        assert item["duration"] == duration
        assert len(item["outcomes"]) == rows
        assert sum(row["weekly_hours"] * 12 for row in item["outcomes"]) == total
        for quarter in range(1, duration + 1):
            assert sum(row["weekly_hours"] for row in item["outcomes"] if row["quarter"] == quarter) == weekly
        assert all(row["result"] and row["source_sheet"] and row["source_row"] >= 3 for row in item["outcomes"])


def test_reimports_and_replacement_preserve_classification_and_other_shift(tmp_path, curriculum_files):
    db = tmp_path / "db.sqlite3"
    original = import_curricula(db, curriculum_files)
    rows = original["competencies"]
    next(row for row in rows if row["key"] == "INGLES")["transversal"] = False
    save_competencies(db, rows)
    catalog = import_curricula(db, curriculum_files + curriculum_files)
    assert len(catalog["competencies"]) == 18
    assert len(catalog["curricula"]) == 2
    assert not next(row for row in catalog["competencies"] if row["key"] == "INGLES")["transversal"]
    name, content = curriculum_files[0]
    workbook = load_workbook(BytesIO(content))
    workbook["Trimestre 1"]["E3"] = 1.5
    output = BytesIO()
    workbook.save(output)
    updated = import_curricula(db, [(name, output.getvalue())])
    assert next(item for item in updated["curricula"] if item["schedule"] == "Mixta") == next(item for item in catalog["curricula"] if item["schedule"] == "Mixta")
    assert not next(row for row in updated["competencies"] if row["key"] == "INGLES")["transversal"]
    assert sum(len(item["outcomes"]) for item in updated["curricula"]) == 89


def workbook_bytes(quarters=(1,), hours=5, competence="  Matemáticas.  "):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for quarter in quarters:
        sheet = workbook.create_sheet(f"Trimestre {quarter}")
        sheet.append(["Competencia", "Resultado", "Horas semanales"])
        sheet.append([competence, "Resultado de aprendizaje", hours])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.mark.parametrize("hours", [-1, None, "abc", "=1+1", True])
def test_bad_hours_rejected(hours):
    with pytest.raises(ValueError, match="horas semanales inválidas"):
        parse_curriculum(workbook_bytes(hours=hours), "Programa - DIURNA.xlsx")


def test_invalid_batch_and_classification_roll_back(tmp_path, curriculum_files):
    db = tmp_path / "db.sqlite3"
    original = import_curricula(db, curriculum_files)
    with pytest.raises(ValueError):
        import_curricula(db, [("Nuevo - DIURNA.xlsx", workbook_bytes()), ("Mal - MIXTA.xlsx", b"invalid")])
    assert load_curricula(db) == original
    rows = [dict(row) for row in original["competencies"]]
    rows[0]["transversal"] = True
    rows[-1]["key"] = "UNKNOWN"
    with pytest.raises(ValueError):
        save_competencies(db, rows)
    assert load_curricula(db) == original
    with pytest.raises(ValueError, match="sin saltos"):
        parse_curriculum(workbook_bytes((1, 3)), "Programa - DIURNA.xlsx")


def test_competencies_deduplicated_across_programs_and_spelling(tmp_path):
    catalog = import_curricula(tmp_path / "db.sqlite3", [("Uno - DIURNO.xlsx", workbook_bytes()),
                                                       ("Dos - MIXTA.xlsx", workbook_bytes(competence="MATEMATICAS"))])
    assert len(catalog["competencies"]) == 1
    assert len(catalog["curricula"]) == 2


def test_actual_hours_shared_plant_and_endings(curriculum_catalog):
    for row in curriculum_catalog["competencies"]:
        row["transversal"] = False
    _, plan = run_plan(curriculum_catalog)
    # Meta 100 = dos continuaciones + dos nuevas. Q1: 112 h/sem; Q2–Q4: 56.
    assert plan["center"]["demanda_total_horas_anuales"] == 3360
    assert plan["center"]["demanda_tecnica_horas_anuales"] == 3360
    assert plan["center"]["demanda_bilinguismo_horas_anuales"] == 0
    assert plan["center"]["demanda_integralidad_horas_anuales"] == 0
    assert [r["Fichas nuevas"] for r in plan["quarterly"]] == [2, 0, 0, 0]
    assert all(r["Capacidad planta (h/sem)"] == 32 for r in plan["technical_quarterly"])
    assert [r["Contratistas requeridos"] for r in plan["technical_quarterly"]] == [2, 1, 1, 1]
    pending = [row for row in plan["curriculum_hours"] if row["Cohorte"].startswith("Ficha")]
    assert {row["Trimestre calendario"] for row in pending} == {1}
    assert {row["Trimestre de formación"] for row in pending} == {7, 9}
    assert sum(row["Horas del trimestre"] for row in pending) == 672


def test_global_classification_changes_components_not_total(curriculum_catalog):
    for row in curriculum_catalog["competencies"]:
        row["transversal"] = row["key"] == "INGLES"
    _, plan = run_plan(curriculum_catalog)
    # Inglés diurno: continuación 3 + nueva T1 [0,3,3,3] = 12.
    # Inglés mixto: continuación 4 + nueva T1 [0,4,0,4] = 12.
    expected = (12 + 12) * 12  # 288 horas, comprobadas contra las hojas originales.
    assert plan["center"]["demanda_bilinguismo_horas_anuales"] == expected
    assert plan["center"]["demanda_total_horas_anuales"] == 3360
    assert plan["center"]["demanda_tecnica_horas_anuales"] == 3360 - expected
    assert all(row["Área"] == "Bilingüismo" for row in plan["curriculum_hours"] if name_key(row["Competencia"]) == "INGLES")


def test_missing_shift_blocks_and_report_is_recomputed(curriculum_catalog):
    catalog = {**curriculum_catalog, "curricula": [c for c in curriculum_catalog["curricula"] if c["schedule"] == "Diurna"]}
    with pytest.raises(ValueError, match="Faltan mallas.*Mixta"):
        run_plan(catalog)
    _, plan = run_plan(curriculum_catalog)
    imported = plan["ficha_import"]
    imported["summary"][0]["Fichas que pasan"] = 500
    instructors = instructors_frame()
    fresh = execute_curriculum_plan(instructors, imported, curriculum_catalog, PlanningRules(intake_weights=(100, 0, 0, 0)),
                                    {"Técnico": 0, "Tecnólogo": 100}, 2027, "instructores.xlsx", "instructores")
    assert fresh["continuing_fichas"] == 2


def test_malla_duration_and_hours_override_nominal_rules(tmp_path):
    catalog = import_curricula(tmp_path / "db.sqlite3", [(PROGRAM + " - DIURNA.xlsx", workbook_bytes((1, 2), 5))])
    frame = ficha_frame(day_age=1).iloc[:1]
    _, plan = run_plan(catalog, frame=frame, target=25)
    assert plan["ficha_import"]["detail"][0]["Duración (trimestres)"] == 2
    assert plan["ficha_import"]["detail"][0]["Trimestre fin estimado"] == 1
    # Cada ficha demanda 5 h, aunque el parámetro general diga 30.
    assert all(row["Horas totales (h/sem)"] == row["Fichas activas"] * 5 for row in plan["calendar"])
    assert all(row["Horas totales del trimestre"] == row["Fichas activas"] * 60 for row in plan["calendar"])


def test_save_reload_export_freezes_catalog(tmp_path, curriculum_catalog, curriculum_files):
    db = tmp_path / "db.sqlite3"
    import_curricula(db, curriculum_files)
    instructors, plan = run_plan(curriculum_catalog)
    save_planning(db, instructors, plan)
    recovered = load_planning(db)
    assert recovered[1]["curriculum_catalog"] == curriculum_catalog
    assert load_curricula(db) == curriculum_catalog
    rows = curriculum_catalog["competencies"]
    rows[0]["transversal"] = not rows[0]["transversal"]
    save_competencies(db, rows)
    assert load_planning(db)[1]["curriculum_catalog"]["competencies"][0]["transversal"] != rows[0]["transversal"]
    exported = load_workbook(BytesIO(export_planning(*recovered)), data_only=True)
    assert {"Mallas curriculares", "Competencias", "Resultados curriculares", "Trazabilidad horas"}.issubset(exported.sheetnames)
    assert exported["Resultados curriculares"].max_row == 90
    assert exported["Competencias"].max_row == 19
    sheet = exported["Trazabilidad horas"]
    column = [c.value for c in sheet[1]].index("Horas del trimestre")
    assert sum(row[column] for row in sheet.iter_rows(min_row=2, values_only=True)) == 3360


def test_real_report_identifies_all_missing_programs(curriculum_catalog):
    frame = parse_fichas_excel(Path(__file__).parents[1] / "data" / "reporteFichas_2026_4.xlsx")
    instructors = parse_instructors_excel(Path(__file__).parent / "fixtures" / "reporteInstructores.xlsx")
    with pytest.raises(ValueError, match="Faltan mallas curriculares") as error:
        prepare_ficha_import(frame, instructors, curriculum_catalog, 2027, "reporte.xlsx", "real", 2026, 4)
    assert "PROGRAMACION DE SOFTWARE" in str(error.value)
    assert "ANALISIS Y DESARROLLO DE SOFTWARE" not in str(error.value)


def test_weeks_and_offer_change_annual_hours_without_mutating_catalog(curriculum_catalog):
    _, plan = run_plan(curriculum_catalog, rules=PlanningRules(weeks_per_quarter=10, intake_weights=(100, 0, 0, 0)))
    assert plan["center"]["demanda_total_horas_anuales"] == 2800
    _, late = run_plan(curriculum_catalog, frame=ficha_frame(7, 9), rules=PlanningRules(intake_weights=(0, 0, 0, 100)), target=50)
    # Ambas continuaciones ya terminaron; dos nuevas en T4, una por jornada.
    assert late["center"]["demanda_total_horas_anuales"] == 672
    assert late["continuing_fichas"] == 0
    assert [row["Fichas nuevas"] for row in late["quarterly"]] == [0, 0, 0, 2]
