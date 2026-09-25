from collections import Counter, defaultdict
from copy import deepcopy
from datetime import date, timedelta
from io import BytesIO
import math
import re

from openpyxl import load_workbook
import pytest

from core.database import load_planning, save_planning
from core.export import export_planning
from core.virtual_competencies import assign_teaching_profiles, competency_rows
from core.virtual_schedule_planner import VirtualRules
from core.virtual_staffing_reports import contract_reports, profile_peak_rows, schedule_activity_rows
from test_virtual_schedules import FIXTURES, configured_catalog, person, plan, real_catalog, save_choices

CODE = "240202501"
PROFILE = "Bilingüismo"
DAY = timedelta(days=1)


def programs(catalog):
    return [{"Programa": item["program"], "Incluir": True,
             "Nivel": "Tecnólogo" if i == 0 else "Técnico", "Peso de oferta": 1}
            for i, item in enumerate(catalog["schedules"])]


def test_all_bilingual_excel_codes_and_texts_are_read_even_without_language_keywords(tmp_path):
    catalog = real_catalog(tmp_path)
    for filename, item in zip(("adso_fases.xlsx", "ciberseguridad_fases.xlsx"), catalog["schedules"]):
        book = load_workbook(FIXTURES / filename, data_only=True)
        source = {}
        for sheet in book:
            for cells in sheet.iter_rows():
                for cell in cells:
                    text = " ".join(str(cell.value or "").split())
                    match = re.match(r"(GA\d+-240202501-AA\d+)", text)
                    if match:
                        source[match[1]] = text
        book.close()
        parsed = [row for row in item["activities"] if row["competency"] == CODE]
        assert len(parsed) == len(source) == 7  # Conteo del archivo de prueba, nunca del cálculo.
        assert {row["activity_code"]: row["activity"] for row in parsed} == source
        assert all(row["teaching_profile"] == PROFILE and row["teaching_type"] == "Transversal" for row in parsed)
        for row in parsed:
            row.update(activity="Resultado reformulado sin indicar idioma", teaching_profile="Transversal general",
                       profile_source="manual", teaching_type="Técnico", classification_source="manual")
        assign_teaching_profiles(parsed)
        assert all(row["teaching_profile"] == PROFILE and row["teaching_type"] == "Transversal" for row in parsed)
    audit = [row for row in schedule_activity_rows(catalog) if row["Perfil"] == PROFILE]
    assert len(audit) == 14
    assert len({row["Programa"] for row in audit}) == 2
    assert {row["Competencia"] for row in audit} == {CODE}


@pytest.mark.parametrize("index,offer,finish,start,until", [
    (0, 0, None, "2027-01-09", "2027-12-08"),
    (1, 0, None, "2027-01-09", "2027-08-16"),
    (0, 1, None, "2027-04-09", "2027-12-31"),
    (1, 1, None, "2027-04-09", "2027-11-15"),
    (0, 0, "2028-04-30", "2027-01-01", "2027-07-08"),
    (1, 0, "2027-06-30", "2027-01-01", "2027-05-16"),
])
def test_bilingual_blocks_respect_each_program_offer_and_continuing_end(tmp_path, index, offer, finish, start, until):
    catalog = real_catalog(tmp_path)
    program = programs(catalog)[index]
    cohorts = [{"Programa": program["Programa"], "Fichas que pasan": 10, "Fecha fin lectiva": finish}] if finish else []
    targets = {"Técnico": 0, "Tecnólogo": 0, program["Nivel"]: 250}
    _, result = plan(catalog, programs=[program], cohorts=cohorts, targets=targets,
                     rules=VirtualRules(intake_weights=tuple(100 if i == offer else 0 for i in range(4))))
    timeline = [row for row in result["activity_hours"] if row["Perfil"] == PROFILE]
    assert min(row["Inicio"] for row in timeline) == start
    assert max(row["Fin"] for row in timeline) == until
    assert all(row["Horas semanales por ficha"] == 2 for row in timeline)
    day = date(2027, 1, 1)
    hours = 0
    while day.year == 2027:
        if day.weekday() < 5:
            expected = start <= day.isoformat() <= until
            active = [row for row in timeline if row["Inicio"] <= day.isoformat() <= row["Fin"]]
            assert len(active) == int(expected)  # Dos AA en un bloque no duplican la ficha.
            hours += 10 * 2 / 5 if expected else 0
        day += DAY
    assert sum(row["Horas instructor en vigencia"] for row in timeline) == pytest.approx(hours)
    assert max(row["Contratistas requeridos"] for row in result["staffing"] if row["Perfil"] == PROFILE) == 1


@pytest.mark.parametrize("fichas,plant_count,capacity,required", [
    (16, 0, 40, 1), (20, 0, 40, 1), (21, 0, 40, 2),
    (16, 1, 40, 0), (17, 1, 40, 1), (37, 1, 40, 2), (16, 0, 30, 2),
])
def test_bilingual_contractors_are_derived_from_hours_not_a_fixed_count(tmp_path, fichas, plant_count, capacity, required):
    _, catalog = configured_catalog(tmp_path)
    activity = catalog["schedules"][0]["activities"][1]
    for n in range(2, 7):
        duplicate = deepcopy(activity)
        duplicate.update(id=f"extra{n}", activity=f"Otro resultado de inglés {n}")
        catalog["schedules"][0]["activities"].append(duplicate)
    _, result = plan(catalog, targets={"Técnico": 0, "Tecnólogo": fichas * 25},
                     plant=[person(kind="Transversal", profile=PROFILE)] if plant_count else [],
                     rules=VirtualRules(weekly_contractor_hours=capacity, intake_weights=(100, 0, 0, 0)))
    high = max((row for row in result["staffing"] if row["Perfil"] == PROFILE), key=lambda r: r["Fichas activas"])
    assert high["Atenciones activas"] == fichas
    assert high["Horas requeridas (h/sem)"] == fichas * 2
    assert high["Contratistas requeridos"] == required
    assert high["Capacidad planta (h/sem)"] == plant_count * 32


@pytest.mark.parametrize("field,value", [("Tipo", "Técnico"), ("Perfil docente", "Transversal general")])
def test_bilingual_competency_cannot_be_lost_or_counted_as_general(tmp_path, field, value):
    path, catalog = configured_catalog(tmp_path)
    before = path.read_bytes()
    choices = competency_rows(catalog)
    next(row for row in choices if row["Competencia"] == CODE)[field] = value
    with pytest.raises(ValueError, match="perfil exclusivo Bilingüismo"):
        save_choices(path, catalog, choices)
    assert path.read_bytes() == before


def test_daily_capacity_ledger_matches_all_transversal_periods_months_and_contracts(tmp_path):
    catalog = real_catalog(tmp_path)
    program_rows = programs(catalog)
    _, result = plan(catalog, programs=program_rows, rules=VirtualRules(), targets={"Técnico": 750, "Tecnólogo": 900},
                     cohorts=[{"Programa": program_rows[0]["Programa"], "Fichas que pasan": 5, "Fecha fin lectiva": "2028-04-30"},
                              {"Programa": program_rows[1]["Programa"], "Fichas que pasan": 3, "Fecha fin lectiva": "2027-06-30"}],
                     plant=[person("123", "Transversal", PROFILE), person("456", "Transversal", "Transversal general")])
    monthly = Counter()
    day = date(2027, 1, 1)
    while day.year == 2027:
        if day.weekday() < 5:
            tasks = defaultdict(set)
            for row in result["activity_hours"]:
                if row["Tipo"] == "Transversal" and row["Inicio"] <= day.isoformat() <= row["Fin"]:
                    for number in range(row["Fichas"]):
                        tasks[row["Perfil"]].add((row["Programa"], row["Cohorte"], number, row["Competencias"]))
            for profile in [PROFILE, "Transversal general", "Cultura física"]:
                row, = [row for row in result["staffing"] if row["Perfil"] == profile and row["Inicio"] <= day.isoformat() <= row["Fin"]]
                units = len(tasks[profile])
                plant_slots = 16 if profile in {PROFILE, "Transversal general"} else 0
                contracts = math.ceil(max(0, units - plant_slots) / 20)
                assert row["Atenciones activas"] == units
                assert row["Horas requeridas (h/sem)"] == units * 2
                assert row["Contratistas requeridos"] == contracts
                current = [r for r in result["contracts"] if r["Perfil"] == profile and r["Inicio"] <= day.isoformat() <= r["Fin"]]
                assert len(current) == len({r["Cupo"] for r in current}) == contracts
                monthly[day.month] += units * 2 / 5
        day += DAY
    for month, expected in monthly.items():
        actual = sum(row["Horas transversales"] for row in result["monthly_fichas"] if row["Mes"] == month)
        assert actual == pytest.approx(expected)
    for row in result["monthly_instructors"]:
        assert row["Horas asignadas"] <= row["Capacidad en horas"] + 1e-9


def test_contract_slot_with_three_periods_is_one_person_and_retains_all_gaps(tmp_path):
    _, catalog = configured_catalog(tmp_path, durations=(7,))
    first = catalog["schedules"][0]["activities"][0]
    first.update(teaching_type="Transversal", teaching_profile="Transversal general", competency="123456789")
    for index, start in enumerate([14, 28], 2):
        another = deepcopy(first)
        another.update(id=f"block{index}", block_id=f"B{index}", start_offset=start)
        catalog["schedules"][0]["activities"].append(another)
    staff, result = plan(catalog, targets={"Técnico": 0, "Tecnólogo": 21 * 25})
    slots, periods = contract_reports(list(reversed(result["contracts"])))
    assert len(slots) == 2
    assert len(periods) == 6
    assert all(row["Número de períodos"] == 3 for row in slots)
    assert all(row["Pausas sin contratación"] == "2027-01-08 a 2027-01-14 · 2027-01-22 a 2027-01-28" for row in slots)
    assert [row["Período del cupo"] for row in periods if row["Cupo"] == 2] == ["1 de 3", "2 de 3", "3 de 3"]
    peaks = profile_peak_rows(result["staffing"], result["rules"])
    assert peaks[0]["Pico simultáneo de contratistas"] == 2
    assert peaks[0]["Máxima carga (h/sem)"] == 42
    assert result["summary"]["pico_contratistas_total"] == 2
    path = tmp_path / "saved.sqlite3"
    save_planning(path, staff, result)
    book = load_workbook(BytesIO(export_planning(*load_planning(path))), read_only=True)
    assert book["Cupos de contratacion"].max_row == 3
    assert book["Periodos por cupo"].max_row == 7
    assert book["Picos por perfil"].max_row == 2
    assert book["Resultados por perfil"].max_row == 4
    book.close()


def test_no_demand_has_no_slots_or_peak():
    assert contract_reports([]) == ([], [])
    assert profile_peak_rows([], {"weekly_contractor_hours": 40}) == []


def test_profiles_with_different_peak_dates_do_not_inflate_simultaneous_total(tmp_path):
    _, catalog = configured_catalog(tmp_path)
    _, result = plan(catalog)
    # El técnico solo se necesita en semanas 1/3; Bilingüismo, en semana 2.
    peaks = profile_peak_rows(result["staffing"], result["rules"])
    assert sum(row["Pico simultáneo de contratistas"] for row in peaks) == 2
    assert result["summary"]["pico_contratistas_total"] == 1


def test_older_saved_report_without_attention_counts_can_still_be_downloaded(tmp_path):
    path, catalog = configured_catalog(tmp_path)
    staff, result = plan(catalog)
    result.pop("workload_model")
    for row in result["staffing"]:
        row.pop("Atenciones activas")
        row.pop("Atenciones cubiertas por planta")
    for item in result["schedule_catalog"]["schedules"]:
        for row in item["activities"]:
            row.pop("teaching_profile", None)
    save_planning(path, staff, result)
    before = path.read_bytes()
    book = load_workbook(BytesIO(export_planning(*load_planning(path))), read_only=True)
    assert {"Cupos de contratacion", "Resultados por perfil", "Picos por perfil"} <= set(book.sheetnames)
    assert book["Picos por perfil"].max_row == 3
    book.close()
    assert path.read_bytes() == before
