from copy import deepcopy
from io import BytesIO

from openpyxl import load_workbook
import pandas as pd
import pytest

from core.config import PlanningRules
from core.curriculum_store import import_curricula, load_curricula, save_competencies
from core.curriculum_planner import execute_curriculum_plan, prepare_ficha_import
from core.database import load_planning, reset_planning_database, save_planning
from core.export import export_planning
from core.planning_modules import planning_database
from core.virtual_planning import execute_virtual_plan
from test_curriculum_automation import malla


@pytest.fixture
def virtual_catalog(tmp_path):
    return import_curricula(tmp_path / "virtual.sqlite3", [("Software - VIRTUAL.xlsx", malla([10, 20, 30]))])


def virtual_run(catalog, *, cohorts=None, plant=None, programs=None, target=100, rules=None):
    return execute_virtual_plan(
        catalog,
        programs if programs is not None else [{"Programa": "Software", "Incluir": True, "Nivel": "Tecnólogo", "Peso de oferta": 1}],
        cohorts if cohorts is not None else [{"Programa": "Software", "Fichas que pasan": 2, "Trimestre al iniciar la vigencia": 2}],
        plant if plant is not None else [{"Área": "Técnica", "Perfil": "Software", "Instructores de planta": 1}],
        rules or PlanningRules(intake_weights=(100, 0, 0, 0)), {"Técnico": 0, "Tecnólogo": target}, 2027)


def test_virtual_manual_counts_hours_and_endings(virtual_catalog):
    staff, plan = virtual_run(virtual_catalog)
    assert len(staff) == 1
    assert plan["center"]["fichas_que_pasan"] == 2
    assert plan["center"]["fichas_nuevas"] == 2
    # T1: 2*20 + 2*10; T2: 2*30 + 2*20; T3: 2*30; T4: 0.
    assert [r["Horas requeridas"] for r in plan["quarterly"]] == [720, 1200, 720, 0]
    assert plan["center"]["demanda_total_horas_anuales"] == 2640
    assert [r["contratistas_totales"] for r in plan["quarterly"]] == [1, 2, 1, 0]
    assert {r["Fecha fin estimada"] for r in plan["ficha_import"]["detail"]} == {"2027-06-30"}
    assert sum(r["Horas asignadas (h/mes)"] for r in plan["monthly_assignments"]) == 2640
    assert {r["Jornada"] for r in plan["distribution"]} == {"Virtual"}


def test_virtual_zero_plant_and_zero_continuations_can_save_and_export(tmp_path, virtual_catalog):
    staff, plan = virtual_run(virtual_catalog, cohorts=[], plant=[], target=50)
    assert staff.empty
    assert plan["center"]["fichas_que_pasan"] == 0
    assert plan["center"]["fichas_nuevas"] == 2
    assert [r["contratistas_totales"] for r in plan["quarterly"]] == [1, 1, 2, 0]
    path = tmp_path / "saved_virtual.sqlite3"
    save_planning(path, staff, plan)
    restored = load_planning(path)
    assert restored[0].empty
    assert restored[1]["virtual_inputs"] == plan["virtual_inputs"]
    book = load_workbook(BytesIO(export_planning(*restored)), read_only=True)
    assert {"Programas virtuales", "Fichas virtuales manuales", "Planta virtual manual"} <= set(book.sheetnames)
    rows = list(book["Resumen planeacion"].values)
    assert rows[1][rows[0].index("Modalidad")] == "Virtual"
    book.close()


def test_virtual_no_demand_or_continuing_above_target(virtual_catalog):
    _, empty = virtual_run(virtual_catalog, cohorts=[], plant=[], target=0)
    assert empty["center"]["demanda_total_horas_anuales"] == 0
    assert empty["summary"]["pico_contratistas_total"] == 0
    _, continuing = virtual_run(virtual_catalog, target=0)
    assert continuing["center"]["fichas_nuevas"] == 0
    assert continuing["center"]["demanda_total_horas_anuales"] == 1200


def test_virtual_and_presencial_use_the_same_calculation(tmp_path, virtual_catalog):
    staff, virtual = virtual_run(virtual_catalog)
    presencial_catalog = import_curricula(tmp_path / "presencial.sqlite3", [("Software - DIURNA.xlsx", malla([10, 20, 30]))])
    fichas = pd.DataFrame([{"Ficha": str(i), "Especialidad": "Software", "Nivel": "Tecnólogo",
                            "Jornada": "Diurna", "Trimestre actual": 1} for i in (1, 2)])
    imported = prepare_ficha_import(fichas, staff, presencial_catalog, 2027, "fichas.xlsx", "digest", 2026, 4)
    presencial = execute_curriculum_plan(staff, imported, presencial_catalog, PlanningRules(intake_weights=(100, 0, 0, 0)),
                                         {"Técnico": 0, "Tecnólogo": 100}, 2027, "planta.xlsx", "digest")
    for field in ("center", "summary", "quarterly", "monthly", "levels", "contract_windows", "contracting_quarterly"):
        assert virtual[field] == presencial[field], field


def test_different_cohort_ages_and_last_quarter_are_preserved(virtual_catalog):
    _, plan = virtual_run(virtual_catalog, target=50, cohorts=[
        {"Programa": "Software", "Fichas que pasan": 1, "Trimestre al iniciar la vigencia": 1},
        {"Programa": "Software", "Fichas que pasan": 1, "Trimestre al iniciar la vigencia": 3},
    ])
    assert [r["Fichas activas"] for r in plan["quarterly"]] == [2, 1, 1, 0]
    assert [r["Horas requeridas"] for r in plan["quarterly"]] == [480, 240, 360, 0]


@pytest.mark.parametrize("field,value", [("Fichas que pasan", -1), ("Fichas que pasan", 1.5),
                                         ("Fichas que pasan", True), ("Trimestre al iniciar la vigencia", 0),
                                         ("Trimestre al iniciar la vigencia", 4), ("Programa", "Sin malla")])
def test_invalid_manual_cohorts_are_rejected(virtual_catalog, field, value):
    row = {"Programa": "Software", "Fichas que pasan": 1, "Trimestre al iniciar la vigencia": 2, field: value}
    with pytest.raises(ValueError):
        virtual_run(virtual_catalog, cohorts=[row])


@pytest.mark.parametrize("value", [-1, 0.5, True, float("inf")])
def test_invalid_plant_counts_are_rejected(virtual_catalog, value):
    with pytest.raises(ValueError, match="Instructores de planta"):
        virtual_run(virtual_catalog, plant=[{"Área": "Técnica", "Perfil": "Software", "Instructores de planta": value}])


def test_manual_weights_and_intake_percentages(tmp_path):
    catalog = import_curricula(tmp_path / "db.sqlite3", [(p + " - VIRTUAL.xlsx", malla([10, 20])) for p in ("Software", "Redes")])
    _, plan = virtual_run(catalog, cohorts=[], plant=[], target=200, programs=[
        {"Programa": "Software", "Incluir": True, "Nivel": "Tecnólogo", "Peso de oferta": 3},
        {"Programa": "Redes", "Incluir": True, "Nivel": "Tecnólogo", "Peso de oferta": 1},
    ], rules=PlanningRules(intake_weights=(50, 25, 25, 0)))
    assert {r["Especialidad"]: r["Fichas nuevas"] for r in plan["distribution"]} == {"Redes": 2, "Software": 6}
    assert [r["Fichas nuevas"] for r in plan["quarterly"]] == [4, 2, 2, 0]
    assert all("Peso de oferta" in r and "Fichas del reporte" not in r for r in plan["offers_by_program"])


def test_same_program_modalities_are_isolated_and_reset_is_local(tmp_path):
    presencial = planning_database(tmp_path / "planeacion.sqlite3", "Presencial")
    virtual = planning_database(presencial, "Virtual")
    assert presencial != virtual
    catalog_p = import_curricula(presencial, [("Software - DIURNA.xlsx", malla([12, 12]))])
    catalog_v = import_curricula(virtual, [("Software - VIRTUAL.xlsx", malla([10, 20, 30]))])
    before = presencial.read_bytes()
    staff, plan = virtual_run(catalog_v)
    save_planning(virtual, staff, plan)
    changed = deepcopy(catalog_v["competencies"])
    changed[0]["transversal"] = True
    save_competencies(virtual, changed)
    assert presencial.read_bytes() == before
    assert load_curricula(presencial) == catalog_p
    reset_planning_database(virtual)
    assert load_planning(virtual) is None
    assert load_curricula(virtual)["curricula"] == []
    assert presencial.read_bytes() == before
