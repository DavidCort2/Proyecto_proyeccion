from io import BytesIO
import math
import sqlite3

from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from core.competency_identity import competency_identity
from core.config import PlanningRules
from core.contracting_periods import contracting_headline, individual_contract_periods
from core.curriculum import curriculum_key, curriculum_lookup, duration_lookup
from core.curriculum_planner import execute_curriculum_plan, prepare_ficha_import
from core.curriculum_store import import_curricula, load_curricula, save_competencies
from core.database import save_planning, load_planning
from core.export import export_planning
from test_curriculum import curriculum_files, curriculum_catalog, run_plan, ficha_frame, instructors_frame, PROGRAM


@pytest.mark.parametrize("alias,canonical", [
    ("APLICACIÓN DE CONOCIMIENTOS DE LAS CIENCIAS NATURALES DE ACUERDO CON SITUACI S DEL CONTEXTO PRODUCTIVO Y SOCIAL", "FÍSICA"),
    ("APLICACIÓN DE CONOCIMIENTOS DE LAS CIENCIAS NATURALES DE ACUERDO CON SITUACIONES DEL CONTEXTO PRODUCTIVO Y SOCIAL - FÍSICA", "FÍSICA"),
    ("APLICACIÓN DE CONOCIMIENTOS DE LA FÍSICA DE ACUERDO CON SITUACIONES DEL CONTEXTO PRODUCTIVO Y SOCIAL", "FÍSICA"),
    ("APLICAR PRÁCTICAS DE PROTECCIÓN AMBIENTAL, SEGURIDAD Y SALUD EN TRABAJO DE ACUERDO CON LAS POLÍTICAS ORGANIZACIONALES Y LA NORMATIVIDAD VIGENTE", "PROTECCIÓN AMBIENTAL, SEGURIDAD Y SALUD EN EL TRABAJO"),
    ("APLICAR PRÁCTICAS DE PROTECCIÓN AMBIENTAL, SEGURIDAD Y SALUD EN EL TRABAJO DE ERDO CON LAS POLÍTICAS ORGANIZACIONALES Y LA NORMATIVIDAD VIGENTE", "PROTECCIÓN AMBIENTAL, SEGURIDAD Y SALUD EN EL TRABAJO"),
    ("RAZONAR CUANTITATIVAMENTE FRENTE A SITUACIONES SUSCEPTIBLES DE SER ABORDADA E MANERA MATEMÁTICA EN CONTEXTOS LABORALES, SOCIALES Y PERSONALES", "MATEMÁTICAS"),
    ("RAZONAR CUANTITATIVAMENTE FRENTE A SITUACIONES SUSCEPTIBLES DE SER ABORDADAS DE MANERA MATEMÁTICA EN CONTEXTOS LABORALES, SOCIALES", "MATEMÁTICAS"),
    ("DESARROLLAR PROCESOS DE COMUNICACIÓN EFICACES Y EFECTIVOS, TENIENDO EN CUE SITUACIONES DE ORDEN SOCIAL, PERSONAL Y PRODUCTIVO", "COMUNICACIÓN"),
    ("UTILIZAR HERRAMIENTAS INFORMÁTICAS DE ACUERDO CON NECESIDADES DE MANEJO DE RMACIÓN", "TIC"),
])
def test_observed_truncated_names_and_aliases_are_one_competency(alias, canonical):
    assert competency_identity(alias) == competency_identity(canonical)


def test_distinct_technical_and_language_skills_are_not_merged():
    assert competency_identity("CONTROL DE CALIDAD DE PROCESOS DE MEDICIÓN")[0] != competency_identity("CONTROL DE LA SEGURIDAD DE LA INFORMACIÓN DIGITAL")[0]
    assert competency_identity("COMPRENDER TEXTOS EN INGLÉS EN FORMA ESCRITA Y AUDITIVA")[0] != competency_identity("PRODUCIR TEXTOS EN INGLÉS EN FORMA ESCRITA Y ORAL")[0]


def malla(hours, competence="Competencia técnica", kind="Técnica"):
    book = Workbook()
    book.remove(book.active)
    for quarter, weekly in enumerate(hours, 1):
        sheet = book.create_sheet(f"Trimestre {quarter}")
        sheet.append(["Competencia", "Resultado", "Horas semanales", "Tipo competencia"])
        sheet.append([competence, f"Aprendizaje trimestre {quarter}", weekly, kind])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_migrates_existing_duplicates_without_losing_hours_or_originals(tmp_path):
    path = tmp_path / "legacy.sqlite3"
    names = ["FISICA", "APLICACIÓN DE CONOCIMIENTOS DE LAS CIENCIAS NATURALES DE ACUERDO CON SITUACIONES DEL CONTEXTO PRODUCTIVO Y SOCIAL",
             "APLICACIÓN DE CONOCIMIENTOS DE LAS CIENCIAS NATURALES DE ACUERDO CON SITUACI S DEL CONTEXTO PRODUCTIVO Y SOCIAL", "INGLES", "Competencia técnica"]
    with sqlite3.connect(path) as db:
        db.executescript("""
            CREATE TABLE curricula(id INTEGER PRIMARY KEY,program TEXT,program_key TEXT,schedule TEXT,duration INTEGER,source_name TEXT,source_digest TEXT,UNIQUE(program_key,schedule));
            CREATE TABLE competencies(id INTEGER PRIMARY KEY,name_key TEXT UNIQUE,name TEXT,transversal INTEGER,transversal_area TEXT);
            CREATE TABLE curriculum_outcomes(curriculum_id INTEGER,competency_id INTEGER,quarter INTEGER,result TEXT,weekly_hours REAL,source_type TEXT,source_sheet TEXT,source_row INTEGER,PRIMARY KEY(curriculum_id,source_sheet,source_row));
            CREATE TABLE execution(payload TEXT);
            INSERT INTO execution VALUES ('ultima ejecucion intacta');
            INSERT INTO curricula VALUES(1,'Programa','PROGRAMA','Diurna',1,'Programa - DIURNA.xlsx','old');
        """)
        for index, name in enumerate(names, 1):
            # La última fue seleccionada a mano aunque el Excel decía técnica.
            db.execute("INSERT INTO competencies VALUES(?,?,?,?,?)", (index, name, name, int(index == 5), "Integralidad"))
            db.execute("INSERT INTO curriculum_outcomes VALUES(1,?,1,?,5,?,'Trimestre 1',?)",
                       (index, f"Resultado {index}", "Integralidad" if index <= 3 else "Bilingüismo" if index == 4 else "Técnica", index + 1))
    catalog = load_curricula(path)
    assert len(catalog["competencies"]) == 3
    assert all(row["transversal"] for row in catalog["competencies"])
    assert sum(row["weekly_hours"] for row in catalog["curricula"][0]["outcomes"]) == 25
    assert [row["source_competency"] for row in catalog["curricula"][0]["outcomes"]] == names
    assert len(catalog["curricula"][0]["outcomes"]) == 5
    assert load_curricula(path) == catalog
    with sqlite3.connect(path) as db:
        assert db.execute("SELECT * FROM execution").fetchone()[0] == "ultima ejecucion intacta"
        assert not db.execute("PRAGMA foreign_key_check").fetchall()


def test_manual_false_and_area_survive_other_mallas_and_reimports(tmp_path):
    path = tmp_path / "db.sqlite3"
    file = ("Uno - DIURNA.xlsx", malla([3], "FISICA", "Integralidad"))
    catalog = import_curricula(path, [file])
    assert catalog["competencies"][0]["transversal"]
    catalog["competencies"][0].update(transversal=False, transversal_area="Bilingüismo")
    save_competencies(path, catalog["competencies"])
    alias = "APLICACIÓN DE CONOCIMIENTOS DE LA FÍSICA DE ACUERDO CON SITUACIONES DEL CONTEXTO PRODUCTIVO Y SOCIAL"
    updated = import_curricula(path, [file, ("Dos - MIXTA.xlsx", malla([4], alias, "Integralidad"))])
    assert len(updated["competencies"]) == 1
    assert updated["competencies"][0]["classification_source"] == "manual"
    assert not updated["competencies"][0]["transversal"]
    assert updated["competencies"][0]["transversal_area"] == "Bilingüismo"


def test_automatic_classification_reconciles_new_excel_types(tmp_path):
    path = tmp_path / "db.sqlite3"
    catalog = import_curricula(path, [("Uno - DIURNA.xlsx", malla([3], "Competencia nueva", "Técnica"))])
    save_competencies(path, catalog["competencies"])  # Guardar sin cambios no congela la propuesta.
    updated = import_curricula(path, [("Uno - DIURNA.xlsx", malla([3], "Competencia nueva", "Integralidad"))])
    assert updated["competencies"][0]["transversal"]
    assert updated["competencies"][0]["classification_source"] == "automatic"


def test_each_month_tracks_cohort_age_and_entry_offer(tmp_path):
    catalog = import_curricula(tmp_path / "db.sqlite3", [(PROGRAM + " - DIURNA.xlsx", malla([10, 20, 30]))])
    _, plan = run_plan(catalog, frame=ficha_frame(day_age=3).iloc[:1], rules=PlanningRules(intake_weights=(0, 100, 0, 0)), target=25)
    assert [row["Horas requeridas"] for row in plan["monthly"]] == [0, 0, 0, 40, 40, 40, 80, 80, 80, 120, 120, 120]
    assert {row["Mes número"] for row in plan["monthly_fichas"]} == set(range(4, 13))
    assert {row["Trimestre de formación"] for row in plan["monthly_fichas"] if row["Mes número"] == 4} == {1}
    assert {row["Trimestre de formación"] for row in plan["monthly_fichas"] if row["Mes número"] == 10} == {3}
    assert plan["center"]["demanda_total_horas_anuales"] == 720


def test_continuing_ficha_only_receives_remaining_months_and_new_ids_are_stable(curriculum_catalog):
    _, plan = run_plan(curriculum_catalog)
    frame = pd.DataFrame(plan["monthly_fichas"])
    for ficha, expected, age in [("001", 120, 7), ("002", 104, 9)]:
        part = frame.loc[frame["Ficha"] == ficha]
        assert part["Mes número"].tolist() == [1, 2, 3]
        assert part["Horas requeridas (h/mes)"].tolist() == [expected] * 3
        assert part["Trimestre de formación"].tolist() == [age] * 3
    new = frame.loc[frame["Tipo ficha"] == "Nueva proyectada"]
    assert new["Ficha"].nunique() == 2
    assert not frame.duplicated(["Ficha", "Mes número"]).any()
    for quarter in plan["quarterly"]:
        months = [row for row in plan["monthly"] if row["Trimestre"] == quarter["Trimestre"]]
        assert sum(row["Horas requeridas"] for row in months) == quarter["Horas requeridas"]
        assert all(row["Contratistas requeridos"] == quarter["contratistas_totales"] for row in months)


def test_shared_plant_capacity_and_current_contractors_are_ignored(tmp_path):
    catalog = import_curricula(tmp_path / "db.sqlite3", [(PROGRAM + " - DIURNA.xlsx", malla([20, 20, 20]))])
    instructors = instructors_frame()
    instructors.loc[len(instructors)] = [PROGRAM, "Técnica", "Contrato técnico", "04", "Contratista", 5.0, False]
    source = prepare_ficha_import(ficha_frame(day_age=3).iloc[:1], instructors, catalog, 2027, "f.xlsx", "x", 2026, 4)
    plan = execute_curriculum_plan(instructors, source, catalog, PlanningRules(intake_weights=(100, 0, 0, 0)),
                                   {"Técnico": 0, "Tecnólogo": 50}, 2027, "i.xlsx", "x")
    january = next(row for row in plan["monthly_staffing"] if row["Área"] == "Técnica" and row["Mes número"] == 1)
    assert january["Horas requeridas (h/mes)"] == 160
    assert january["Capacidad planta (h/mes)"] == 128  # 32*4, compartidas entre las dos fichas.
    assert january["Contratistas requeridos"] == 1
    assert "Contratistas adicionales" not in january
    assert "Contratistas actuales" not in january
    people = [row for row in plan["monthly_instructors"] if row["Mes número"] == 1 and row["Área"] == "Técnica"]
    assert next(row for row in people if row["Tipo"] == "Planta")["Fichas atendidas"] == 2
    assert all(row["Instructor"] != "Contrato técnico" for row in people)
    assert sum(row["Tipo"] == "Contratista proyectado" for row in people) == 1
    assert all(row["Horas asignadas (h/mes)"] <= row["Capacidad (h/mes)"] for row in plan["monthly_instructors"])
    assert sum(row["Horas asignadas (h/mes)"] for row in people) == 160
    assert sum(row["Horas asignadas (h/mes)"] for row in plan["monthly_assignments"]) == plan["center"]["demanda_total_horas_anuales"]


def test_free_capacity_of_other_profile_cannot_cover_a_deficit(tmp_path):
    catalog = import_curricula(tmp_path / "db.sqlite3", [(PROGRAM + " - DIURNA.xlsx", malla([20, 20, 20]))])
    _, plan = run_plan(catalog, frame=ficha_frame(day_age=3).iloc[:1], target=50)
    # Planta libre en Bilingüismo e Integralidad no cubre el déficit técnico.
    assert [row["Contratistas requeridos"] for row in plan["monthly"]] == [1] * 9 + [0] * 3
    assert plan["summary"]["pico_contratistas_total"] == 1
    assert all(row["Horas asignadas (h/mes)"] == 0 for row in plan["monthly_instructors"] if row["Área"] != "Técnica")


def test_fractional_weeks_conserve_hours_and_assignments(curriculum_catalog):
    _, plan = run_plan(curriculum_catalog, rules=PlanningRules(weeks_per_quarter=10, intake_weights=(100, 0, 0, 0)))
    assert math.fsum(row["Horas requeridas"] for row in plan["monthly"]) == pytest.approx(2800)
    assert math.fsum(row["Horas asignadas (h/mes)"] for row in plan["monthly_assignments"]) == pytest.approx(2800)
    assert all(row["Horas asignadas (h/mes)"] <= row["Capacidad (h/mes)"] + 1e-8 for row in plan["monthly_instructors"])


def test_monthly_snapshot_and_excel_match_saved_execution(tmp_path, curriculum_catalog):
    instructors, plan = run_plan(curriculum_catalog)
    path = tmp_path / "db.sqlite3"
    save_planning(path, instructors, plan)
    recovered = load_planning(path)
    assert recovered[1]["monthly_fichas"] == plan["monthly_fichas"]
    assert recovered[1]["contract_windows"] == plan["contract_windows"]
    book = load_workbook(BytesIO(export_planning(*recovered)), data_only=True)
    assert book.sheetnames[:2] == ["Contratacion requerida", "Contratistas y fechas"]
    assert dict(zip(next(book["Contratacion requerida"].values), list(book["Contratacion requerida"].values)[1])) == contracting_headline(plan)
    sheet = book["Contratistas y fechas"]
    exported_periods = [dict(zip(next(sheet.values), row)) for row in list(sheet.values)[1:]]
    assert exported_periods == individual_contract_periods(plan)
    assert {"Resumen mensual", "Horas mensuales por ficha", "Dotacion mensual por perfil", "Capacidad mensual instructores", "Asignacion mensual de horas", "Supuestos mensuales"}.issubset(book.sheetnames)
    assert {"Periodos de contratacion", "Contratacion por trimestre", "Picos por perfil", "Excesos y reducciones", "Criterio de contratacion"}.issubset(book.sheetnames)
    sheet = book["Horas mensuales por ficha"]
    headings = [cell.value for cell in sheet[1]]
    assert {"Archivo malla", "Duración (trimestres)", "Fecha fin estimada"}.issubset(headings)
    assert "Criterio de duraciones" in book.sheetnames
    column = [cell.value for cell in sheet[1]].index("Horas requeridas (h/mes)")
    assert sum(row[column] for row in sheet.iter_rows(min_row=2, values_only=True)) == 3360


@pytest.mark.parametrize("duration", [2, 7, 10, 12])
def test_program_abbreviation_does_not_create_a_shift_or_impose_a_duration(duration):
    item = {"program": "DESARROLLO Y ADAPTACION DE ORTESIS Y PROTESIS", "schedule": "Diurna", "duration": duration}
    catalog = {"curricula": [item]}
    key = curriculum_key("DESARROLLO Y ADAPTACION DE PROTESIS Y ORTESIS", "P&O-tarde")
    assert key == curriculum_key(item["program"], "DIURNO")
    assert curriculum_lookup(catalog)[key] == item
    assert duration_lookup(catalog)[key] == duration
    assert len(curriculum_lookup(catalog)) == 1
