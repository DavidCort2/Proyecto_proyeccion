"""Lectura de mallas: identidad por programa/jornada y horas semanales por resultado."""
from hashlib import sha256
from io import BytesIO
import math
from pathlib import Path
import re

from openpyxl import load_workbook

from core.excel_parser import normalize_text
from core.competency_identity import competency_identity


def name_key(value):
    return " ".join(re.sub(r"[^\w\s]", " ", normalize_text(value)).split())


def schedule_name(value):
    value = normalize_text(value)
    compact = value.replace(" ", "")
    if compact.startswith(("O&P", "P&O", "DIURNAO&P", "DIURNOO&P")):
        return "Diurna O&P"
    if value in {"DIURNA", "DIURNO"} or value.startswith(("DIURNA-", "DIURNO-")):
        return "Diurna"
    if value in {"MIXTA", "MIXTO"}:
        return "Mixta"
    raise ValueError(f"Jornada no reconocida: {value}.")


def curriculum_key(program, schedule):
    program = name_key(program)
    # El mismo programa aparece con el orden de los dos dispositivos invertido.
    if program == "DESARROLLO Y ADAPTACION DE ORTESIS Y PROTESIS":
        program = "DESARROLLO Y ADAPTACION DE PROTESIS Y ORTESIS"
    return program, schedule_name(schedule)


def parse_curriculum(content: bytes, filename: str) -> dict:
    """Lee horas y tipo del Excel; conserva el título original de cada resultado."""
    stem = Path(filename).stem
    match = re.fullmatch(r"(.+?)\s+[-–—]\s+(.+)", stem)
    if not match:
        raise ValueError(f"{filename}: use el nombre PROGRAMA - JORNADA.xlsx.")
    program, schedule = match.group(1).strip(), schedule_name(match.group(2))
    try:
        book = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"{filename}: no es un Excel .xlsx válido.") from exc
    outcomes, quarters = [], set()
    try:
        for sheet in book:
            match = re.fullmatch(r"TRIMESTRE\s+(\d+)", normalize_text(sheet.title))
            if not match:
                # Se permiten portadas, pero no se descartan hojas de datos desconocidas.
                if any("COMPETENCIA" == normalize_text(cell) for row in sheet.iter_rows(values_only=True) for cell in row):
                    raise ValueError(f"{filename}, {sheet.title}: nombre de hoja sin trimestre. Use Trimestre N.")
                continue
            quarter = int(match.group(1))
            if quarter < 1 or quarter > 40 or quarter in quarters:
                raise ValueError(f"{filename}: trimestre duplicado o inválido ({quarter}).")
            quarters.add(quarter)
            header = None
            count = 0
            for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                cells = [normalize_text(value) for value in row]
                required = {"COMPETENCIA", "RESULTADO", "HORAS SEMANALES"}
                if required.issubset(cells):
                    header = {name: cells.index(name) for name in required}
                    header["TIPO COMPETENCIA"] = cells.index("TIPO COMPETENCIA") if "TIPO COMPETENCIA" in cells else None
                    continue
                if header is None or not any(value is not None and str(value).strip() for value in row):
                    continue
                def field(name):
                    index = header[name]
                    return row[index] if index is not None and index < len(row) else None
                competence, result, raw_hours = field("COMPETENCIA"), field("RESULTADO"), field("HORAS SEMANALES")
                location = f"{filename}, {sheet.title}, fila {number}"
                if not name_key(competence) or result is None or not str(result).strip():
                    raise ValueError(f"{location}: faltan competencia o resultado.")
                try:
                    if isinstance(raw_hours, bool):
                        raise ValueError()
                    hours = float(str(raw_hours).strip().replace(",", "."))
                    if not math.isfinite(hours) or hours < 0:
                        raise ValueError()
                except (TypeError, ValueError):
                    raise ValueError(f"{location}: horas semanales inválidas o fórmula sin valor calculado.") from None
                identity, display = competency_identity(competence)
                outcomes.append({"quarter": quarter, "competency_key": identity,
                                 "competency": display, "source_competency": " ".join(str(competence).split()), "result": str(result).strip(),
                                 "weekly_hours": hours, "source_type": str(field("TIPO COMPETENCIA") or ""),
                                 "source_sheet": sheet.title, "source_row": number})
                count += 1
            if not count:
                raise ValueError(f"{filename}, {sheet.title}: faltan resultados y horas semanales.")
        if not quarters or quarters != set(range(1, max(quarters) + 1)):
            raise ValueError(f"{filename}: las hojas deben cubrir todos los trimestres desde el 1, sin saltos.")
    finally:
        book.close()
    return {"program": program, "program_key": curriculum_key(program, schedule)[0], "schedule": schedule,
            "duration": max(quarters), "source_name": Path(filename).name,
            "source_digest": sha256(content).hexdigest(), "outcomes": outcomes}


def suggested_transversal_area(outcome):
    label = normalize_text(outcome["source_type"])
    name = name_key(outcome["competency"])
    return "Bilingüismo" if "BILING" in label or re.search(r"\b(INGLES|INGLESA|ENGLISH)\b", name) else "Integralidad"


def catalog_digest(catalog):
    import json
    return sha256(json.dumps(catalog, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def duration_lookup(catalog):
    return {key: item["duration"] for key, item in curriculum_lookup(catalog).items()}


def curriculum_lookup(catalog):
    lookup = {curriculum_key(item["program"], item["schedule"]): item for item in catalog["curricula"]}
    # O&P es diurna de diez trimestres. Solo admite esta equivalencia si la
    # malla diurna confirma diez trimestres; nunca toma una diurna regular de 7.
    for (program, schedule), item in list(lookup.items()):
        if schedule == "Diurna" and item["duration"] == 10:
            lookup.setdefault((program, "Diurna O&P"), item)
    return lookup
