"""Extrae una plantilla de programa por competencias y duración, sin fechas de ficha."""
from calendar import monthrange
from datetime import date, datetime
from fractions import Fraction
from hashlib import sha256
from io import BytesIO
import json
import math
from pathlib import Path
import re

from openpyxl import load_workbook

from core.curriculum import name_key
from core.excel_parser import normalize_text


def is_lective_activity(row):
    """La etapa productiva se conserva como origen, pero no genera carga docente."""
    return normalize_text(row["phase"]) != "ETAPA PRODUCTIVA"


def lective_activities(schedule):
    return [row for row in schedule["activities"] if is_lective_activity(row)]


def read_date(value, label):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(value).strip(), pattern).date()
        except ValueError:
            pass
    raise ValueError(f"{label}: indique una fecha válida (día/mes/año).")


def finite_hours(value, label):
    try:
        hours = float(value)
        if isinstance(value, bool) or not math.isfinite(hours) or hours < 0:
            raise ValueError()
    except (TypeError, ValueError):
        raise ValueError(f"{label}: ingrese horas no negativas; use 0 explícito si no requiere instructor.") from None
    return hours


def _source_number(value, label):
    normalized = re.sub(r"([.,])\s+(?=\d)", r"\1", normalize_text(value))
    match = re.fullmatch(r"\s*(\d+(?:[.,]\d+)?)\s*(?:HORAS?|MESES?|MES|DIAS?|SEMANAS?)?\s*", normalized)
    if not match:
        raise ValueError(f"{label}: no se pudo leer el valor {value!r}.")
    return float(match[1].replace(",", "."))


def require_program_template(item):
    if item.get("schema_version") != 3:
        raise ValueError(f"{item['program']}: vuelva a cargar el cronograma para leer la duración general del programa; las fechas de la ficha de origen no se usan para planear.")


def duration_sum(*values):
    """Suma decimales del Excel sin introducir fracciones binarias en los límites."""
    return float(sum((Fraction(str(value)) for value in values), Fraction(0)))


def duration_boundary(start, offset, unit):
    """Desplazamiento acumulado: meses calendario, fracción prorrateada y días completos."""
    from datetime import timedelta
    value = Fraction(str(offset))
    if unit in {"días", "semanas"}:
        return start + timedelta(days=math.ceil(value * (7 if unit == "semanas" else 1)))
    if unit != "meses":
        raise ValueError("Unidad de duración del programa no reconocida.")
    whole = math.floor(value)

    def anniversary(months):
        year, month = divmod(start.year * 12 + start.month - 1 + months, 12)
        return date(year, month + 1, min(start.day, monthrange(year, month + 1)[1]))

    beginning, following = anniversary(whole), anniversary(whole + 1)
    return beginning + timedelta(days=math.ceil((value - whole) * (following - beginning).days))


def _phases(sheet, column, first, last):
    """Reconstruye rótulos partidos al convertir un cronograma PDF a Excel."""
    groups = []
    for row in range(first, last + 1):
        raw = sheet.cell(row, column).value
        if raw is None:
            continue
        compact = re.sub(r"\s+", "", normalize_text(raw))
        if compact.startswith(("ETAPA", "TOTAL")):
            continue
        if compact.startswith(("FAS", "INDUC")):
            groups.append([row, compact])
        elif groups:
            groups[-1][1] += compact
        else:
            raise ValueError(f"{sheet.title}, fila {row}: no se reconoce el nombre de la fase.")
    result = {}
    for index, (start, text) in enumerate(groups):
        if text.startswith("INDUC"):
            label = "Inducción"
        else:
            match = re.fullmatch(r"FASE(\d+)(.+)", text)
            if not match:
                raise ValueError(f"{sheet.title}, fila {start}: complete el nombre de la fase ({text}).")
            names = {"ANALISIS": "Análisis", "PLANEACION": "Planeación", "EJECUCION": "Ejecución", "EVALUACION": "Evaluación"}
            title = match[2].strip(".:-")
            label = f"Fase {int(match[1])} · {names.get(title, title.capitalize())}"
        end = groups[index + 1][0] if index + 1 < len(groups) else last + 1
        result.update({row: label for row in range(start, end)})
    return result


def parse_virtual_schedule(content, filename):
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Los cronogramas virtuales se cargan exclusivamente en Excel (.xlsx).")
    try:
        book = load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:
        raise ValueError(f"{filename}: no es un Excel válido.") from exc
    program, project, activities, blocks, summaries, warnings = None, "", [], [], [], []
    duration_unit = None
    try:
        for sheet in book:
            header = None
            for row in sheet.iter_rows():
                cells = {normalize_text(cell.value): cell.column for cell in row if cell.value is not None}
                if {"FASES", "ACTIVIDADES DE APRENDIZAJE"} <= set(cells):
                    header = row[0].row, cells
                    break
            if header is None:
                continue
            header_row, columns = header
            if "ACTIVIDADES DE PROYECTO" in columns:
                columns["ACTIVIDADES DEL PROYECTO"] = columns["ACTIVIDADES DE PROYECTO"]
            needed = ["ACTIVIDADES DEL PROYECTO"]
            if any(label not in columns for label in needed):
                raise ValueError(f"{sheet.title}: falta la columna de actividades del proyecto.")
            duration_columns = [(label, unit) for label, unit in (
                ("TIEMPO DE DURACION ESTIMADO MESES", "meses"),
                ("TIEMPO DE DURACION ESTIMADO DIAS", "días"),
                ("TIEMPO DE DURACION ESTIMADO SEMANAS", "semanas")) if label in columns]
            if not duration_columns and "TIEMPO DE DURACION ESTIMADO" in columns:
                detected = set()
                for cells in sheet.iter_rows(min_row=header_row + 1):
                    value_text = normalize_text(cells[columns["TIEMPO DE DURACION ESTIMADO"] - 1].value)
                    for pattern, candidate in [(r"\bMES(?:ES)?\b", "meses"), (r"\bDIAS?\b", "días"), (r"\bSEMANAS?\b", "semanas")]:
                        if re.search(pattern, value_text):
                            detected.add(candidate)
                if len(detected) == 1:
                    duration_columns = [("TIEMPO DE DURACION ESTIMADO", detected.pop())]
            if len(duration_columns) != 1:
                raise ValueError(f"{sheet.title}: indique una columna de duración estimada en meses, días o semanas; no se deduce de las fechas de una ficha.")
            duration_column, unit = duration_columns[0]
            if duration_unit is not None and duration_unit != unit:
                raise ValueError("Use la misma unidad de duración en las hojas del programa.")
            duration_unit = unit
            for row in sheet.iter_rows(max_row=header_row - 1, values_only=True):
                for raw in row:
                    if not isinstance(raw, str):
                        continue
                    program_match = re.search(r"Nombre\s+de\s+programa\s*:\s*(.*?)(?=\s*Nombre\s+del\s+Proyecto\s*:|$)", raw, re.I | re.S)
                    if program_match:
                        name = " ".join(program_match[1].split())
                        if program and name_key(program) != name_key(name):
                            raise ValueError("Cargue un cronograma por programa en cada archivo.")
                        program = name
                    project_match = re.search(r"Nombre\s+del\s+Proyecto\s*:\s*(.*)", raw, re.I | re.S)
                    if project_match:
                        project = " ".join(project_match[1].split())
            # Solo propaga celdas realmente combinadas; los vacíos de una
            # continuación de página se tratan después dentro del mismo bloque.
            merged = {}
            for span in sheet.merged_cells.ranges:
                for row in range(span.min_row, span.max_row + 1):
                    for col in range(span.min_col, span.max_col + 1):
                        merged[(row, col)] = (span.min_row, span.min_col)

            def value(row, label):
                if label not in columns:
                    return None, None
                col = columns[label]
                anchor = merged.get((row, col), (row, col))
                return sheet.cell(*anchor).value, anchor

            phases = _phases(sheet, columns["FASES"], header_row + 1, sheet.max_row)
            current = None
            counted_hours = set()
            for row in range(header_row + 1, sheet.max_row + 1):
                location = f"{sheet.title}, fila {row}"
                unanchored = False
                phase_raw = normalize_text(value(row, "FASES")[0])
                if phase_raw.startswith("TOTAL") and "ETAPA LECTIVA" in phase_raw:
                    phase_raw = "ETAPA LECTIVA"
                duration_raw, _ = value(row, duration_column)
                hours_raw, hours_anchor = value(row, "TIEMPO DE DURACION ESTIMADO HORAS")
                learning = sheet.cell(row, columns["ACTIVIDADES DE APRENDIZAJE"]).value
                if phase_raw in {"ETAPA LECTIVA", "ETAPA PRODUCTIVA"}:
                    hours = _source_number(hours_raw, location) if hours_raw is not None else None
                    duration = _source_number(duration_raw, location)
                    summaries.append({"stage": phase_raw, "duration": duration, "duration_unit": unit, "hours": hours})
                    if phase_raw == "ETAPA LECTIVA":
                        continue  # Resumen, no otra actividad: no duplica horas.
                    learning = "Seguimiento de etapa productiva"
                    phase = "Etapa productiva"
                    project_activity = "Etapa productiva"
                else:
                    if learning is None or not str(learning).strip():
                        continue
                    phase = phases.get(row)
                    if phase is None:
                        raise ValueError(f"{location}: actividad sin fase.")
                    project_activity = value(row, "ACTIVIDADES DEL PROYECTO")[0]
                    if (project_activity is None or normalize_text(project_activity) in {"AP", "AP-"}) and duration_raw is None and current and current["phase"] == phase:
                        project_activity = current["project_activity"]
                        unanchored = True
                block_key = (phase, str(project_activity))
                if current is None or current["identity"] != block_key:
                    duration = _source_number(duration_raw, f"{location}, duración")
                    if duration <= 0:
                        raise ValueError(f"{location}: la duración estimada del bloque debe ser mayor que cero.")
                    current = {"id": f"B{len(blocks) + 1}", "identity": block_key, "phase": phase,
                               "project_activity": str(project_activity or ""),
                               "start_offset": duration_sum(*(block["duration"] for block in blocks)),
                               "duration": duration, "duration_unit": unit,
                               "source_hours": 0.0 if hours_raw is not None else None, "source_sheet": sheet.title, "source_row": row}
                    blocks.append(current)
                elif duration_raw is not None and _source_number(duration_raw, location) != current["duration"]:
                    raise ValueError(f"{location}: hay duraciones distintas dentro de la misma actividad del proyecto.")
                anchor_key = (sheet.title, hours_anchor)
                if hours_raw is not None and anchor_key not in counted_hours:
                    current["source_hours"] = (current["source_hours"] or 0.0) + _source_number(hours_raw, location)
                    counted_hours.add(anchor_key)
                raw_text = " ".join(str(learning).split())
                matches = list(re.finditer(r"GA\s*\d+\s*-\s*(\d{6,12})\s*-\s*AA\s*\d+", raw_text, flags=re.I))
                # Una celda puede contener varias actividades o competencias.
                # Cada código tiene clasificación y carga propia; las horas de
                # referencia continúan contándose una sola vez por bloque.
                if len(matches) > 1:
                    warnings.append(f"{location}: contiene {len(matches)} códigos de actividad; se extraen por separado sus competencias.")
                entries = [(re.sub(r"\s+", "", match[0].upper()), match[1],
                            raw_text[match.start() if i else 0:matches[i + 1].start() if i + 1 < len(matches) else len(raw_text)].strip())
                           for i, match in enumerate(matches)] or [(f"{phase} · {raw_text}", phase, raw_text)]
                if not matches and phase not in {"Inducción", "Etapa productiva"}:
                    # Un fragmento sin código no permite inferir con seguridad a
                    # qué competencia pertenece. Se conserva sin inventar otra.
                    current.setdefault("source_notes", []).append({"source_row": row, "text": raw_text})
                    warnings.append(f"{location}: fragmento sin código conservado en las notas del bloque; no crea una competencia ni carga adicional.")
                    continue
                for identity, competency, activity_text in entries:
                    if any(old["activity_code"] == identity and old["block_id"] == current["id"]
                           and name_key(old["activity"]) == name_key(activity_text) for old in activities):
                        warnings.append(f"{location}: actividad repetida {identity}; se cuenta una sola vez.")
                        continue
                    activities.append({"id": sha256((name_key(program) + "|" + identity + "|" + current["id"]).encode()).hexdigest()[:24],
                                       "block_id": current["id"], "phase": phase, "competency": competency,
                                       "activity_code": identity, "activity": activity_text, "source_text": raw_text,
                                       "start_offset": current["start_offset"], "duration": current["duration"], "duration_unit": unit,
                                       "teaching_type": "Técnico", "classification_source": "automatic", "instructor_hours": None,
                                       "unanchored": unanchored,
                                       "source_sheet": sheet.title, "source_row": row})
        if not program or not activities:
            raise ValueError("No se encontró un cronograma con Nombre de programa, Fases, Actividades de Aprendizaje y duración estimada.")
        if not any(is_lective_activity(row) for row in activities):
            raise ValueError("El cronograma debe contener actividades de etapa lectiva para planear sus horas.")
        # En los Excel convertidos desde PDF hay filas de un GA antes del
        # encabezado de su bloque. La pertenencia se resuelve con los otros
        # registros del mismo GA que sí tienen proyecto y duración, no por
        # proximidad ni por un número de fila propio de un programa.
        anchored = {}
        for activity in activities:
            prefix = re.match(r"GA\d+", activity["activity_code"])
            if prefix and not activity["unanchored"]:
                anchored.setdefault((activity["source_sheet"], prefix[0]), set()).add(activity["block_id"])
        indexed_blocks = {block["id"]: block for block in blocks}
        for activity in activities:
            if activity.pop("unanchored"):
                prefix = re.match(r"GA\d+", activity["activity_code"])
                candidates = anchored.get((activity["source_sheet"], prefix[0] if prefix else ""), set())
                if len(candidates) != 1:
                    raise ValueError(f"{activity['source_sheet']}, fila {activity['source_row']}: no se puede asociar de forma única {activity['activity_code']} con un bloque; complete su proyecto y duración.")
                block = indexed_blocks[next(iter(candidates))]
                activity.update(block_id=block["id"], phase=block["phase"], start_offset=block["start_offset"], duration=block["duration"])
                activity["id"] = sha256((name_key(program) + "|" + activity["activity_code"] + "|" + block["id"]).encode()).hexdigest()[:24]
                warnings.append(f"{activity['source_sheet']}, fila {activity['source_row']}: {prefix[0]} se vincula a {block['project_activity']} ({block['phase']}) por su código; no añade otra duración.")
        if any(block["id"] not in {row["block_id"] for row in activities} for block in blocks):
            raise ValueError("Hay un bloque de duración sin competencias identificables. Complete sus códigos en el cronograma.")
        if len({item["id"] for item in activities}) != len(activities):
            raise ValueError("Hay actividades con el mismo código repetido. Revise el cronograma para evitar duplicarlas.")
        lective = math.fsum(row["source_hours"] or 0 for row in blocks if row["phase"] not in {"Inducción", "Etapa productiva"})
        lective_duration = duration_sum(*(row["duration"] for row in blocks if is_lective_activity(row)))
        adjustments = json.loads((Path(__file__).resolve().parents[1] / "config" / "virtual_programs.json").read_text(encoding="utf-8"))["duration_adjustments"]
        adjustment = next((row for row in adjustments if name_key(row["program"]) == name_key(program)), None)
        declared = next((row["duration"] for row in summaries if row["stage"] == "ETAPA LECTIVA"), None)
        applied = []
        if adjustment and declared is not None and declared > lective_duration:
            candidates = [row for row in blocks if normalize_text(adjustment["remainder_phase"]) in normalize_text(row["phase"])]
            if not candidates:
                raise ValueError("No se encontró la fase de evaluación para incorporar el tiempo lectivo confirmado.")
            last = candidates[-1]
            extra = duration_sum(declared, -lective_duration)
            last["duration"] = duration_sum(last["duration"], extra)
            offset = 0.0
            for block in blocks:
                block["start_offset"] = offset
                for activity in activities:
                    if activity["block_id"] == block["id"]:
                        activity.update(start_offset=offset, duration=block["duration"])
                offset = duration_sum(offset, block["duration"])
            applied.append({"phase": last["phase"], "added_duration": extra, "duration_unit": duration_unit, "source": adjustment["source"]})
            warnings.append(f"Se incorporan {extra:g} {duration_unit} a {last['phase']} para completar los {declared:g} declarados, según la aclaración del usuario.")
            lective_duration = declared
        for summary in summaries:
            if summary["stage"] == "ETAPA LECTIVA" and summary["hours"] is not None and not math.isclose(lective, summary["hours"]):
                warnings.append(f"Los bloques lectivos suman {lective:g} h; el resumen declara {summary['hours']:g} h. Diferencia: {summary['hours'] - lective:g} h. No se distribuye ni se corrige automáticamente.")
            if summary["stage"] == "ETAPA LECTIVA" and not math.isclose(lective_duration, summary["duration"]):
                warnings.append(f"Los bloques lectivos, incluida la inducción, suman {lective_duration:g} {unit}; el total del archivo declara {summary['duration']:g}. La proyección respeta las duraciones de los bloques, sin añadir tiempo no asignado.")
        for block in blocks:
            block.pop("identity")
        from core.virtual_competencies import assign_teaching_profiles, suggest_classifications
        suggest_classifications(activities)
        assign_teaching_profiles(activities)
        return {"schema_version": 3, "program": program, "program_key": name_key(program), "project": project,
                "source_name": filename, "source_digest": sha256(content).hexdigest(),
                "duration": duration_sum(*(block["duration"] for block in blocks)), "duration_unit": duration_unit,
                "lective_duration": lective_duration,
                "duration_adjustments": applied,
                "blocks": blocks, "activities": activities,
                "summaries": summaries, "warnings": list(dict.fromkeys(warnings))}
    finally:
        book.close()
